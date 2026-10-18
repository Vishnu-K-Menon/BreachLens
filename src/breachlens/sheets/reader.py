"""Read xlsx / xls / csv into raw grids (SPREADSHEET_TRACK.md §1).

Every sheet becomes a `RawSheet`: typed cell values by 1-based (row, column) with Excel's
own row numbers (CSV: physical line numbers, so comment lines count), merged ranges,
hidden rows/columns, number formats and, for CSV, the detected encoding and delimiter.
Nothing is guessed about types here; CSV values stay strings.

- `.xlsx`: openpyxl (merged ranges, hidden rows, number formats); files above
  `large_file_rows` rows fall back to pandas + calamine (values only).
- `.xls`: xlrd with formatting info (merged cells, hidden rows), dates converted.
- `.csv`: charset-normalizer for the encoding, csv.Sniffer on the first 64 KB for the
  delimiter, leading `#` comment lines skipped.
"""

from __future__ import annotations

import csv
import io
import logging
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

import olefile
import xlrd
from charset_normalizer import from_bytes
from openpyxl import load_workbook

from breachlens.sheets.settings import ReaderCfg

logger = logging.getLogger(__name__)

Cell = str | int | float | datetime | date | None


@dataclass
class RawSheet:
    doc_id: str
    path: str
    file_kind: str
    sheet: str
    rows: dict[int, list[Cell]]  # excel row -> values (index 0 = column 1)
    n_cols: int
    merged: list[tuple[int, int, int, int]] = field(default_factory=list)  # r1, c1, r2, c2
    hidden_rows: set[int] = field(default_factory=set)
    hidden_cols: set[int] = field(default_factory=set)
    formats: dict[tuple[int, int], str] = field(default_factory=dict)
    encoding: str | None = None
    delimiter: str | None = None
    comment_lines: int = 0
    doc_date: str | None = None


def _trim(rows: dict[int, list[Cell]]) -> tuple[dict[int, list[Cell]], int]:
    n_cols = 0
    for vals in rows.values():
        for j in range(len(vals), 0, -1):
            if vals[j - 1] not in (None, ""):
                n_cols = max(n_cols, j)
                break
    return {r: (v + [None] * n_cols)[:n_cols] for r, v in rows.items()}, n_cols


def read_xlsx(path: Path, doc_id: str, rel: str, cfg: ReaderCfg) -> list[RawSheet]:
    wb = load_workbook(path, data_only=True)
    created = wb.properties.created
    out = []
    for ws in wb.worksheets:
        if ws.max_row > cfg.large_file_rows:
            out.append(_read_xlsx_calamine(path, ws.title, doc_id, rel))
            continue
        rows: dict[int, list[Cell]] = {}
        formats: dict[tuple[int, int], str] = {}
        for r, row in enumerate(ws.iter_rows(), start=1):
            vals: list[Cell] = []
            for c, cell in enumerate(row, start=1):
                vals.append(cell.value)
                if cell.value is not None and cell.number_format != "General":
                    formats[(r, c)] = cell.number_format
            rows[r] = vals
        rows, n_cols = _trim(rows)
        sheet = RawSheet(doc_id, rel, "XLSX", ws.title, rows, n_cols, formats=formats)
        sheet.merged = [
            (m.min_row, m.min_col, m.max_row, m.max_col) for m in ws.merged_cells.ranges
        ]
        sheet.hidden_rows = {r for r, d in ws.row_dimensions.items() if d.hidden}
        sheet.hidden_cols = {c for c, d in ws.column_dimensions.items() if d.hidden}
        sheet.doc_date = created.date().isoformat() if created else None
        out.append(sheet)
    return out


def _read_xlsx_calamine(path: Path, sheet: str, doc_id: str, rel: str) -> RawSheet:
    import pandas as pd

    df = pd.read_excel(path, sheet_name=sheet, header=None, dtype=str, engine="calamine")
    rows = {
        i + 1: [None if pd.isna(v) else v for v in r]
        for i, r in enumerate(df.itertuples(index=False))
    }
    rows, n_cols = _trim(rows)
    logger.info("%s/%s: %d rows read with calamine", doc_id, sheet, len(rows))
    return RawSheet(doc_id, rel, "XLSX", sheet, rows, n_cols)


def xls_created(path: Path) -> str | None:
    """Creation date from the OLE SummaryInformation stream (xlrd does not read it)."""
    try:
        with olefile.OleFileIO(str(path)) as ole:
            meta = ole.get_metadata()
    except OSError:
        return None
    return meta.create_time.date().isoformat() if meta.create_time else None


def read_xls(path: Path, doc_id: str, rel: str) -> list[RawSheet]:
    book = xlrd.open_workbook(path, formatting_info=True)
    created = xls_created(path)
    out = []
    for sh in book.sheets():
        rows: dict[int, list[Cell]] = {}
        for r in range(sh.nrows):
            vals: list[Cell] = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    vals.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    vals.append(None)
                elif cell.ctype == xlrd.XL_CELL_ERROR:
                    vals.append(xlrd.error_text_from_code.get(cell.value, "#ERR"))
                else:
                    vals.append(cell.value)
            rows[r + 1] = vals
        rows, n_cols = _trim(rows)
        sheet = RawSheet(doc_id, rel, "XLS", sh.name, rows, n_cols)
        sheet.merged = [(r1 + 1, c1 + 1, r2, c2) for r1, r2, c1, c2 in sh.merged_cells]
        sheet.hidden_rows = {r + 1 for r, info in sh.rowinfo_map.items() if info.hidden}
        sheet.hidden_cols = {c + 1 for c, info in sh.colinfo_map.items() if info.hidden}
        sheet.doc_date = created
        out.append(sheet)
    return out


def detect_encoding(data: bytes, candidates: list[str]) -> str:
    """First candidate that decodes strictly, else charset-normalizer's best guess."""
    for enc in candidates:
        try:
            text = data.decode(enc)
        except UnicodeDecodeError:
            continue
        # cp1252 leaves 5 bytes undefined; C1 control characters mean a wrong guess.
        if not any(0x80 <= ord(ch) <= 0x9F for ch in text):
            return enc
    best = from_bytes(data).best()
    return best.encoding if best else "latin-1"


def read_csv(path: Path, doc_id: str, rel: str, cfg: ReaderCfg) -> list[RawSheet]:
    data = path.read_bytes()
    encoding = detect_encoding(data, cfg.encodings)
    text = data.decode(encoding)
    lines = text.splitlines()
    skip = 0
    while skip < len(lines) and lines[skip].lstrip().startswith(cfg.comment_prefix):
        skip += 1
    body = "\n".join(lines[skip:])
    try:
        delimiter = (
            csv.Sniffer()
            .sniff(body[: cfg.csv_sniff_bytes], delimiters="".join(cfg.csv_delimiters))
            .delimiter
        )
    except csv.Error:
        delimiter = ","
    rows: dict[int, list[Cell]] = {}
    for i, rec in enumerate(csv.reader(io.StringIO(body), delimiter=delimiter), start=skip + 1):
        rows[i] = [v if v != "" else None for v in rec]
    rows, n_cols = _trim(rows)
    return [
        RawSheet(
            doc_id,
            rel,
            "CSV",
            "",
            rows,
            n_cols,
            encoding=encoding,
            delimiter=delimiter,
            comment_lines=skip,
        )
    ]


def read_file(path: Path, doc_id: str, rel: str, kind: str, cfg: ReaderCfg) -> list[RawSheet]:
    if kind == "XLSX":
        return read_xlsx(path, doc_id, rel, cfg)
    if kind == "XLS":
        return read_xls(path, doc_id, rel)
    if kind == "CSV":
        return read_csv(path, doc_id, rel, cfg)
    raise ValueError(f"not a spreadsheet: {kind}")


def as_text(v: Any) -> str:
    """Cell value as a display string (dates ISO, integral floats without .0)."""
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == datetime.min.time() else v.isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer() and abs(v) < 1e15:
        return str(int(v))
    return str(v)
