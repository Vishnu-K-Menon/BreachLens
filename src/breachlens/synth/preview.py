"""Human-review previews (M1 checkpoint): one PDF per type, a scanned example and every
messy spreadsheet, as PNGs plus a markdown summary, written to `data/output/previews/`.

Spreadsheet "screenshots" come from LibreOffice (xlsx -> PDF -> PNG) on a print-scaled
copy, so the real corpus files are never modified.
"""

from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path

import pymupdf as fitz
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter

from breachlens.synth.casting import cast_pdfs
from breachlens.synth.people import build_population, check_reuse_cap, read_pool
from breachlens.synth.records import Rendered
from breachlens.synth.render_pdf import render_pdf
from breachlens.synth.render_sheets import render_all_sheets, soffice_path
from breachlens.synth.scanify import scanify
from breachlens.synth.settings import load_synth_settings


def pdf_to_png(pdf: Path, png: Path, dpi: int = 110, pages: int = 1) -> list[Path]:
    outs = []
    with fitz.open(pdf) as doc:
        for i, page in enumerate(doc):
            if i >= pages:
                break
            target = png if i == 0 else png.with_name(f"{png.stem}_p{i + 1}.png")
            page.get_pixmap(dpi=dpi).save(target)
            outs.append(target)
    return outs


def _soffice(args: list[str], tmp: str) -> None:
    exe = soffice_path()
    if exe is None:
        raise RuntimeError("LibreOffice not found: set SOFFICE_PATH")
    profile = Path(tmp, "lo_profile").as_uri()
    subprocess.run(
        [exe, f"-env:UserInstallation={profile}", "--headless", *args],
        check=True,
        capture_output=True,
        timeout=180,
    )


def sheet_to_png(src: Path, png: Path) -> list[Path]:
    """Landscape, fit-to-width print of each sheet of a copy, rendered with LibreOffice."""
    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp, src.name)
        shutil.copyfile(src, work)
        if work.suffix == ".xls":
            _soffice(["--convert-to", "xlsx", "--outdir", tmp, str(work)], tmp)
            work = work.with_suffix(".xlsx")
        wb = load_workbook(work)
        for ws in wb.worksheets:
            # Readable widths for the screenshot only (the corpus file keeps Excel defaults).
            for idx, col in enumerate(
                ws.iter_cols(min_row=1, max_row=min(ws.max_row, 200)), start=1
            ):
                cells = [c for c in col if not isinstance(c, MergedCell) and c.value is not None]
                longest = max((len(str(c.value)) for c in cells), default=0)
                is_date = any(c.is_date for c in cells)
                width = 11 if is_date and longest < 11 else min(max(6, longest + 1), 34)
                ws.column_dimensions[get_column_letter(idx)].width = width
            ws.page_setup.orientation = "landscape"
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.print_options.gridLines = True
            ws.print_options.headings = True
        printable = Path(tmp, "print_" + work.name)
        wb.save(printable)
        _soffice(["--convert-to", "pdf", "--outdir", tmp, str(printable)], tmp)
        return pdf_to_png(printable.with_suffix(".pdf"), png, dpi=100, pages=2)


def _csv_head(path: Path, n: int = 15) -> list[str]:
    return path.read_bytes().decode("latin-1").splitlines()[:n]


def build_previews(root: Path, out_dir: Path, seed: int | None = None) -> list[Path]:
    settings = load_synth_settings(root / "config", seed)
    cfg = settings.config
    pool = read_pool(root / "data" / "reference" / "address_pool.csv")
    pop = build_population(settings, pool)
    check_reuse_cap(pop, cfg.max_unrelated_households_per_address)
    if out_dir.exists():
        shutil.rmtree(out_dir)
    docs = out_dir / "docs"
    plans = cast_pdfs(pop, settings)
    chosen = []
    for doc_type in cfg.pdf.counts:
        chosen.append(next(p for p in plans if p.doc_type == doc_type and p.variant == "A"))
    scan_plan = next(p for p in plans if p.doc_type == "HEALTH_CLAIM" and p.variant == "B")
    rendered = Rendered()
    files: list[Path] = []
    pdf_notes: list[str] = []
    for plan in [*chosen, scan_plan]:
        plan.scan = plan is scan_plan
        r = render_pdf(plan, pop, docs, cfg.seed)
        rendered.extend(r)
        pdf = docs / plan.doc_file
        if plan.scan:
            scanify(pdf, pdf, cfg.pdf.scan, seed=cfg.seed)
        label = f"{plan.doc_type.lower()}_{plan.variant}{'_scan' if plan.scan else ''}"
        files += pdf_to_png(pdf, out_dir / f"pdf_{label}.png")
        cases = sorted(
            set().union(*(m.planted for m in r.mentions))
            | set().union(*(e.planted for e in r.exclusions))
        )
        who = ", ".join(f"{m.role} {m.person.pid if m.person else 'ORG'}" for m in r.mentions)
        pdf_notes.append(
            f"| `pdf_{label}.png` | {plan.doc_type} | {plan.variant} | {plan.doc_date} | "
            f"{'yes' if plan.scan else 'no'} | {who} | {', '.join(cases) or '-'} |"
        )
    sheets = render_all_sheets(pop, settings, docs)
    lines = [
        "# M1 preview - SYNTHETIC - NOT REAL",
        "",
        f"Seed {cfg.seed}. Population: {len(pop.people)} people, {len(pop.orgs)} organizations.",
        "",
        "## PDFs (one per type + one scanned example)",
        "",
        "| preview | type | variant | doc_date | scan | mentions | planted cases |",
        "|---|---|---|---|---|---|---|",
        *pdf_notes,
        "",
        "## Spreadsheets",
    ]
    for doc in sheets.docs:
        path = docs / doc.doc_file
        stem = Path(doc.doc_file).name.replace(".", "_")
        n_m = sum(1 for m in sheets.mentions if m.doc_file == doc.doc_file and m.person is not None)
        lines += [
            "",
            f"### `{doc.doc_file}` ({doc.file_kind}, doc_date {doc.doc_date or 'none'})",
            "",
            f"Person mentions: {n_m}. Planted cases: {', '.join(sorted(doc.planted))}.",
            "",
        ]
        lines += [f"- {mess}" for mess in doc.messes]
        if doc.file_kind == "CSV":
            lines += [
                "",
                "Encoding: latin-1, delimiter `;`. First 15 raw lines:",
                "",
                "```",
                *_csv_head(path),
                "```",
            ]
        else:
            pngs = sheet_to_png(path, out_dir / f"sheet_{stem}.png")
            files += pngs
            lines += ["", "Screenshot(s): " + ", ".join(f"`{p.name}`" for p in pngs)]
    summary = out_dir / "PREVIEW.md"
    summary.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return [summary, *files]
