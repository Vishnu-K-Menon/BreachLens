"""Render the 8 messy corpus spreadsheets (SYNTHETIC_DATA.md §Spreadsheets).

Files are written with openpyxl so the damage is real: numbers stored as numbers, real
merged ranges, hidden rows, number formats. `legacy_export.xls` is written as xlsx and
converted by LibreOffice. Each xlsx carries `SYNTHETIC - NOT REAL` as an A1 comment and
in its document properties, with the export date (doc_date, ADR-029) as `created`;
the CSV starts with a `# SYNTHETIC - NOT REAL` line.

Row numbers in the returned mentions are 1-based Excel rows (CSV: physical lines).
"""

from __future__ import annotations

import os
import random
import re
import shutil
import subprocess
import tempfile
import zipfile
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet

from breachlens.synth.people import BASE, DEPENDENT, EXTRA_GUARDIAN, STUDENT, Person, Population
from breachlens.synth.records import (
    DocRecord,
    Exclusion,
    Mention,
    Rendered,
    address_fields,
    first_last,
    fmt_ssn,
    last_comma_first,
    mdy,
    name_fields,
    up,
)
from breachlens.synth.settings import SynthSettings

NOTICE = "SYNTHETIC - NOT REAL"
EXCEL_EPOCH = date(1899, 12, 30)
DATE_FMT = "mm/dd/yyyy"
FIXED_ZIP_TIME = (1980, 1, 1, 0, 0, 0)


# ---- shared helpers -------------------------------------------------------------------
def excel_serial(d: date) -> int:
    return (d - EXCEL_EPOCH).days


def _props(wb: Workbook, export: date) -> datetime:
    stamp = datetime(export.year, export.month, export.day, 9, 0, 0)
    wb.properties.creator = "BreachLens synthetic generator"
    wb.properties.lastModifiedBy = "BreachLens synthetic generator"
    wb.properties.title = NOTICE
    wb.properties.subject = NOTICE
    wb.properties.description = f"{NOTICE}. Synthetic test export dated {export.isoformat()}."
    wb.properties.created = stamp
    wb.properties.modified = stamp
    return stamp


def _notice(ws: Worksheet, cell: str = "A1") -> None:
    ws[cell].comment = Comment(NOTICE, "BreachLens")


def save_xlsx(wb: Workbook, path: Path, export: date) -> None:
    """Save, then rewrite the zip with fixed timestamps and the export date as 'modified'.

    openpyxl stamps `modified` with the wall clock and zip entries with the current time;
    both are pinned so the same input gives the same bytes.
    """
    stamp = _props(wb, export).strftime("%Y-%m-%dT%H:%M:%SZ")
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    with zipfile.ZipFile(tmp) as zin, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "docProps/core.xml":
                text = data.decode("utf-8")
                text = re.sub(
                    r"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                    rf"\g<1>{stamp}\g<2>",
                    text,
                )
                text = re.sub(
                    r"(<dcterms:created[^>]*>)[^<]*(</dcterms:created>)",
                    rf"\g<1>{stamp}\g<2>",
                    text,
                )
                data = text.encode("utf-8")
            info = zipfile.ZipInfo(item.filename, date_time=FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            zout.writestr(info, data)
    tmp.unlink()


def _header(ws: Worksheet, row: int, names: list[str]) -> None:
    for j, name in enumerate(names, start=1):
        ws.cell(row=row, column=j, value=name).font = Font(bold=True)


def _mention(
    doc_file: str, sheet: str | None, row: int, role: str, p: Person | None, **kw: object
) -> Mention:
    return Mention(
        doc_file=doc_file,
        source_type="SPREADSHEET",
        role=role,
        person=p,
        sheet=sheet,
        row=row,
        **kw,
    )  # type: ignore[arg-type]


def _full_name(p: Person, first: str | None = None, last: str | None = None) -> str:
    """'First M. Last Suffix' in mixed case, as spreadsheets usually hold it."""
    parts = [first or p.first, f"{p.middle[0]}." if p.middle else "", last or p.last, p.suffix]
    return " ".join(x for x in parts if x)


def _street_inline(p: Person) -> str:
    """Line 1 with the unit inline (P09): 'APT 6F' is printed as 'APT6F'."""
    a = p.address
    if not a.line2:
        return a.line1
    unit = a.line2.replace("APT ", "APT") if a.line2.startswith("APT ") else a.line2
    return f"{a.line1} {unit}"


def _unit_tag(p: Person) -> set[str]:
    return {"P09"} if p.address.line2 else set()


def _cols(sheet: str, labels: list[tuple[str, str]]) -> list[tuple[str, int, str, str]]:
    return [(sheet, i, h, lab) for i, (h, lab) in enumerate(labels, start=1)]


# ---- hr_roster.xlsx ---------------------------------------------------------------------
def render_hr_roster(pop: Population, s: SynthSettings, docs: Path, rng: random.Random) -> Rendered:
    name, sheet = "hr_roster.xlsx", "Roster"
    doc_file = f"sheets/{name}"
    export = s.config.sheets.export_dates[name]
    employer = next(o for o in pop.orgs.values() if o.role == "employer")
    vendor = next(o for o in pop.orgs.values() if o.role == "vendor")
    people = pop.by_cast(BASE) + [p for p in pop.people.values() if p.home_sheet == name]
    rng.shuffle(people)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws["A1"] = employer.name
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Employee roster export - {export.isoformat()}"
    ws["A3"] = f"{NOTICE} - test data only"
    ws.merge_cells("A1:O1")
    _notice(ws)
    groups = [("Employee", "A4:G4"), ("Home Address", "H4:K4"), ("Contact", "L4:M4")]
    for text, rng_ in groups:
        ws[rng_.split(":")[0]] = text
        ws.merge_cells(rng_)
        ws[rng_.split(":")[0]].alignment = Alignment(horizontal="center")
        ws[rng_.split(":")[0]].font = Font(bold=True)
    for col, text in (("N", "Hire Date"), ("O", "Department")):
        ws[f"{col}4"] = text
        ws[f"{col}4"].font = Font(bold=True)
        ws.merge_cells(f"{col}4:{col}5")
    sub = [
        "ID",
        "First",
        "Middle",
        "Last",
        "Suffix",
        "SSN",
        "DOB",
        "Street",
        "City",
        "State",
        "ZIP",
        "Phone",
        "Email",
    ]
    _header(ws, 5, sub)
    depts = ["Operations", "Finance", "Warehouse", "Sales", "IT", "Customer Service", "Fleet"]
    out = Rendered()
    spaced = set(rng.sample([p.pid for p in people], 4))
    row = 6
    row_of: dict[str, int] = {}
    org_at = rng.randint(10, 10 + len(people) - 5)
    dup_pid = rng.choice([p.pid for p in pop.by_cast(BASE)])

    def write_person(p: Person, r: int, tags: set[str]) -> Mention:
        n = name_fields(p)
        first, last = p.first, p.last
        if p.pid in spaced:  # P26
            first, last = f"  {first} ", f"{last}  "
            tags = tags | {"P26"}
        a = address_fields(p.address)
        values = [
            p.employee_id,
            first,
            p.middle,
            last,
            p.suffix,
            fmt_ssn(p.ssn),
            p.dob,
            _street_inline(p),
            a["city"].title(),
            a["state"],
            int(a["zip5"]),
            p.phone,
            p.email,
            p.hire_date,
            rng.choice(depts),
        ]
        for j, val in enumerate(values, start=1):
            c = ws.cell(row=r, column=j, value=val)
            if isinstance(val, date):
                c.number_format = DATE_FMT
        if a["zip5"].startswith("0"):
            tags = tags | {"P31"}
        if p.pid == pop.roles.get("ssn_area_666"):
            tags = tags | {"P13"}
        if p.pid == pop.roles.get("suffix_jr"):
            tags = tags | {"P02"}
        m = _mention(doc_file, sheet, r, "EMPLOYEE", p, planted=tags | _unit_tag(p))
        m.fields = {
            **n,
            **a,
            "dob": p.dob.isoformat(),
            "ssn": p.ssn,
            "employee_id": p.employee_id,
            "phone": p.phone,
            "email": p.email,
        }
        m.printed = {
            "name": f"{first}|{p.middle}|{last}|{p.suffix}",
            "street": _street_inline(p),
            "zip": str(int(a["zip5"])),
        }
        return m

    for i, p in enumerate(people):
        if i == org_at - 6:  # P21: an organization in the person columns
            ws.cell(row=row, column=1, value="C" + str(rng.randint(10000, 99999)))
            ws.cell(row=row, column=4, value=vendor.name)
            ws.cell(row=row, column=15, value="Contractor")
            out.exclusions.append(
                Exclusion(doc_file, "ORG", vendor.name, sheet=sheet, row=row, planted={"P21"})
            )
            mo = _mention(
                doc_file,
                sheet,
                row,
                "EMPLOYEE",
                None,
                org=vendor,
                entity_type="ORG",
                planted={"P21"},
            )
            mo.printed = {"name": vendor.name}
            out.mentions.append(mo)
            row += 1
        out.mentions.append(write_person(p, row, set()))
        row_of[p.pid] = row
        row += 1
        if p.pid == dup_pid:  # P42: exact duplicate row directly below
            for j in range(1, 16):
                src = ws.cell(row=row - 1, column=j)
                c = ws.cell(row=row, column=j, value=src.value)
                c.number_format = src.number_format
            dm = _mention(doc_file, sheet, row, "EMPLOYEE", p, planted={"P42"})
            first = out.mentions[-1]
            dm.fields, dm.printed = dict(first.fields), dict(first.printed)
            dm.fields["duplicate_of_row"] = str(row - 1)
            first.planted.add("P42")
            out.mentions.append(dm)
            row += 1
    ws.cell(row=row, column=1, value="Total employees")
    ws.cell(row=row, column=2, value=sum(1 for m in out.mentions if m.person is not None))
    ws.cell(row=row, column=1).font = Font(bold=True)
    hidden = rng.sample(
        [m for m in out.mentions if m.person is not None and "P42" not in m.planted], 3
    )
    for m in hidden:  # P29
        ws.row_dimensions[m.row].hidden = True
        m.planted.add("P29")
    save_xlsx(wb, docs / doc_file, export)
    labels = [
        ("Employee ID", "EMPLOYEE_ID"),
        ("Employee First", "FIRST_NAME"),
        ("Employee Middle", "MIDDLE_NAME"),
        ("Employee Last", "LAST_NAME"),
        ("Employee Suffix", "SUFFIX"),
        ("Employee SSN", "SSN"),
        ("Employee DOB", "DOB"),
        ("Home Address Street", "ADDRESS_LINE1"),
        ("Home Address City", "CITY"),
        ("Home Address State", "STATE"),
        ("Home Address ZIP", "ZIP"),
        ("Contact Phone", "PHONE"),
        ("Contact Email", "EMAIL"),
        ("Hire Date", "OTHER"),
        ("Department", "OTHER"),
    ]
    doc = DocRecord(
        doc_file,
        "XLSX",
        "SPREADSHEET",
        export,
        planted={"P27", "P28", "P29", "P30", "P31", "P42"},
        columns=_cols(sheet, labels),
    )
    doc.messes = [
        "rows 1-3: title/banner rows above the header (P27)",
        "rows 4-5: merged two-row header, e.g. 'Employee' over 'SSN' (P28)",
        f"hidden rows {sorted(m.row for m in hidden)} (P29)",
        f"row {row}: 'Total employees' summary row (P30)",
        "ZIP stored as a number: NJ/MA ZIPs lost their leading zero (P31)",
        f"row {row_of[dup_pid] + 1}: exact duplicate of row {row_of[dup_pid]} (P42)",
        "an organization (contractor) in the Last column (P21)",
        "leading/trailing spaces in 4 names (P26); units inline in Street, e.g. APT6F (P09)",
    ]
    out.docs.append(doc)
    return out


# ---- payroll_export.csv -----------------------------------------------------------------
def render_payroll(pop: Population, s: SynthSettings, docs: Path, rng: random.Random) -> Rendered:
    name = "payroll_export.csv"
    doc_file = f"sheets/{name}"
    header = [
        "Emp No",
        "Employee Name",
        "SS#",
        "Pay Date",
        "Gross Pay",
        "Net Pay",
        "Bank Routing",
        "Account No",
        "Dept",
        "Street",
        "City",
        "State",
        "Zip",
    ]
    lines = [f"# {NOTICE}", ";".join(header)]
    base = sorted(pop.by_cast(BASE), key=lambda p: p.employee_id)
    jr = pop.roles["suffix_jr"]
    w2_pids = {
        pop.roles[k] for k in ("mover", "nickname", "suffix_jr", "star", "surname_0", "twin_a")
    }
    out = Rendered()
    na_rows = set(rng.sample(range(len(base) * len(s.config.sheets.payroll_pay_dates)), 6))
    jj_done = False
    k = 0
    depts = ["OPS", "FIN", "WHS", "SLS", "IT", "CS"]
    for pay in s.config.sheets.payroll_pay_dates:
        for p in base:
            tags: set[str] = set()
            suffix = p.suffix
            if p.pid == jr and not jj_done:  # P03: doubled-letter suffix typo
                suffix, jj_done = "JJ", True
                tags.add("P03")
            nm = name_fields(p, suffix=suffix)
            head = f"{p.last} {suffix}".strip()
            tail = f"{p.first} {p.middle[:1]}".strip()
            emp_name = f"{head}, {tail}"
            ssn_num = str(int(p.ssn))  # SSN exported as a number: leading zeros lost
            if len(ssn_num) < 9:
                tags.add("P44")
            gross = rng.randint(1500, 6500) + rng.randint(0, 99) / 100
            net = gross * 0.74
            routing, acct = (p.routing, p.bank_account) if p.bank_account else ("", "")
            dept = rng.choice(depts)
            if k in na_rows:  # P43
                dept = "#N/A"
                if acct:
                    acct = "#N/A"
                tags.add("P43")
            a = address_fields(p.address)
            row = [
                p.employee_id,
                emp_name,
                ssn_num,
                pay.strftime("%m/%d/%Y"),
                f"{gross:.2f}".replace(".", ","),
                f"{net:.2f}".replace(".", ","),
                routing,
                acct,
                dept,
                a["line1"] + (f" {a['line2']}" if a["line2"] else ""),
                a["city"].title(),
                a["state"],
                a["zip5"],
            ]
            lines.append(";".join(row))
            if p.pid in w2_pids:
                tags.add("P41")
            if any(ord(ch) > 127 for ch in emp_name):
                tags.add("P35")
            m = _mention(doc_file, None, len(lines), "EMPLOYEE", p, planted=tags)
            m.fields = {**nm, **a, "ssn": p.ssn, "employee_id": p.employee_id}
            if routing and acct and acct != "#N/A":
                m.fields |= {"routing": routing, "bank_account": acct}
            m.printed = {"name": emp_name, "ssn": ssn_num}
            out.mentions.append(m)
            k += 1
    path = docs / doc_file
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(("\r\n".join(lines) + "\r\n").encode("latin-1"))
    labels = [
        ("Emp No", "EMPLOYEE_ID"),
        ("Employee Name", "FULL_NAME"),
        ("SS#", "SSN"),
        ("Pay Date", "OTHER"),
        ("Gross Pay", "OTHER"),
        ("Net Pay", "OTHER"),
        ("Bank Routing", "ROUTING"),
        ("Account No", "BANK_ACCOUNT"),
        ("Dept", "OTHER"),
        ("Street", "ADDRESS_LINE1"),
        ("City", "CITY"),
        ("State", "STATE"),
        ("Zip", "ZIP"),
    ]
    doc = DocRecord(
        doc_file,
        "CSV",
        "SPREADSHEET",
        None,
        planted={"P35", "P43", "P44"},
        columns=_cols("", labels),
    )
    doc.messes = [
        "Latin-1 encoding with accented names, ';' delimiter, decimal commas (P35)",
        "line 1 is a '# SYNTHETIC - NOT REAL' comment; header on line 2",
        "header 'SS#' for the SSN column (P35)",
        "SSN exported as a number: SSNs with a leading zero have 8 digits (P44)",
        "'#N/A' in Dept / Account No cells (P43); 'JJ' suffix typo in one name (P03)",
        "two pay periods per employee (two rows per person)",
    ]
    out.docs.append(doc)
    return out


# ---- patient_list.xlsx ------------------------------------------------------------------
def render_patients(pop: Population, s: SynthSettings, docs: Path, rng: random.Random) -> Rendered:
    name, sheet = "patient_list.xlsx", "Patients"
    doc_file = f"sheets/{name}"
    export = s.config.sheets.export_dates[name]
    insurer = next(o for o in pop.orgs.values() if o.role == "insurer")
    must = [pop.role("mover"), pop.role("star")]
    base = must + rng.sample([p for p in pop.by_cast(BASE) if p not in must], 18)
    people = base + pop.by_cast(DEPENDENT) + pop.by_cast(STUDENT)
    people += [p for p in pop.people.values() if p.home_sheet == name]
    rng.shuffle(people)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    header = [
        "MRN",
        "Patient Name",
        "DOB",
        "SSN",
        "Member ID",
        "Insurer",
        "Address",
        "City",
        "State",
        "ZIP",
        "Phone",
        "Notes",
    ]
    _header(ws, 1, header)
    _notice(ws)
    formats = ["date", "serial", "text_mdy", "text_iso", "text_long", "text_mdy2"]
    note_rows = sorted(rng.sample(range(len(people)), 6))
    benign = [
        "Follow-up in 6 weeks.",
        "Prefers morning appointments.",
        "Allergy list reviewed.",
        "",
        "",
        "Requested paper statements.",
    ]
    out = Rendered()
    for i, p in enumerate(people):
        r = i + 2
        tags: set[str] = set()
        fmt = formats[i % len(formats)] if i < len(formats) else rng.choice(formats)
        dob_val: object = {
            "date": p.dob,
            "serial": excel_serial(p.dob),
            "text_mdy": mdy(p.dob),
            "text_iso": p.dob.isoformat(),
            "text_long": f"{p.dob.day} {p.dob.strftime('%b')} {p.dob.year}",
            "text_mdy2": mdy(p.dob, two_digit_year=True),
        }[fmt]
        if fmt != "date":
            tags.add("P32")
        note = rng.choice(benign)
        fields_extra: dict[str, str] = {}
        if i in note_rows:  # P38: PII in prose
            if note_rows.index(i) % 2 == 0:  # alternate SSN and MRN prose
                note = f"Pt called about a bill; verified SSN {fmt_ssn(p.ssn)} on file."
                fields_extra["notes_pii"] = "SSN"
            else:
                note = f"Chart merged from prior clinic, MRN {p.mrn} confirmed with patient."
                fields_extra["notes_pii"] = "MRN"
            tags.add("P38")
        a = address_fields(p.address)
        values = [
            p.mrn,
            _full_name(p),
            dob_val,
            fmt_ssn(p.ssn),
            p.member_id,
            insurer.name.title(),
            _street_inline(p),
            a["city"].title(),
            a["state"],
            a["zip5"],
            p.phone,
            note,
        ]
        for j, val in enumerate(values, start=1):
            c = ws.cell(row=r, column=j, value=val)
            if isinstance(val, date):
                c.number_format = DATE_FMT
        if p.pid == pop.roles["mover"]:
            tags.add("P17")
        if p.pid == pop.roles["star"]:
            tags.add("P24")
        m = _mention(doc_file, sheet, r, "PATIENT", p, planted=tags | _unit_tag(p))
        m.fields = {
            **name_fields(p),
            **a,
            "dob": p.dob.isoformat(),
            "ssn": p.ssn,
            "mrn": p.mrn,
            "member_id": p.member_id,
            "insurer": insurer.name,
            "phone": p.phone,
            **fields_extra,
        }
        m.printed = {"name": _full_name(p), "dob": str(dob_val), "dob_format": fmt}
        out.mentions.append(m)
    notes = wb.create_sheet("Notes")
    for i, line in enumerate(
        [
            "Export notes",
            "Generated by the clinic's EHR export tool.",
            NOTICE,
            "Contact the records office with questions about this export.",
        ],
        start=1,
    ):
        notes.cell(row=i, column=1, value=line)
    save_xlsx(wb, docs / doc_file, export)
    labels = [
        ("MRN", "MRN"),
        ("Patient Name", "FULL_NAME"),
        ("DOB", "DOB"),
        ("SSN", "SSN"),
        ("Member ID", "MEMBER_ID"),
        ("Insurer", "ORG_NAME"),
        ("Address", "ADDRESS_LINE1"),
        ("City", "CITY"),
        ("State", "STATE"),
        ("ZIP", "ZIP"),
        ("Phone", "PHONE"),
        ("Notes", "FREE_TEXT"),
    ]
    doc = DocRecord(
        doc_file,
        "XLSX",
        "SPREADSHEET",
        export,
        planted={"P32", "P38"},
        columns=_cols(sheet, labels),
    )
    doc.messes = [
        "2 sheets: 'Patients' (table) and 'Notes' (free text only, should be skipped)",
        "DOB column mixes real date cells, Excel serial numbers (e.g. 32874), MM/DD/YYYY, ISO, "
        "'4 Mar 1990' and 2-digit-year text (P32)",
        "Notes column contains SSNs / MRNs in prose on 6 rows (P38)",
        "full name in one column ('First M. Last')",
    ]
    out.docs.append(doc)
    return out


# ---- benefits_enrollment.xlsx -----------------------------------------------------------
def render_benefits(pop: Population, s: SynthSettings, docs: Path, rng: random.Random) -> Rendered:
    name, sheet = "benefits_enrollment.xlsx", "Enrollment"
    doc_file = f"sheets/{name}"
    export = s.config.sheets.export_dates[name]
    deps: dict[str, list[Person]] = {}
    for d in pop.by_cast(DEPENDENT):
        deps.setdefault(d.subscriber, []).append(d)
    subs = [pop.people[pid] for pid in deps]
    others = [p for p in pop.by_cast(BASE) if p.pid not in deps]
    jr = pop.role("suffix_jr")
    pick = subs + ([jr] if jr not in subs else [])
    pick += rng.sample([p for p in others if p not in pick], 34 - len(pick))
    pick += [p for p in pop.people.values() if p.home_sheet == name]
    rng.shuffle(pick)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    header = [
        "Employee ID",
        "First Name",
        "Last Name",
        "Suffix",
        "DOB",
        "SSN",
        "Plan",
        "Group No",
        "Member ID",
        "Address",
        "City",
        "State",
        "ZIP",
    ]
    for k in range(1, 4):
        header += [f"Dep{k}_First", f"Dep{k}_Last", f"Dep{k}_DOB", f"Dep{k}_Rel"]
    _header(ws, 1, header)
    _notice(ws)
    out = Rendered()
    plans = ["PPO Standard", "HMO Plus", "HDHP + HSA"]
    for i, p in enumerate(pick):
        r = i + 2
        tags: set[str] = set()
        suffix = p.suffix
        if p.pid == jr.pid:  # P02: "Junior" spelled out
            suffix = "Junior"
            tags.add("P02")
        a = address_fields(p.address)
        values: list[object] = [
            p.employee_id,
            p.first,
            p.last,
            suffix,
            p.dob,
            fmt_ssn(p.ssn),
            rng.choice(plans),
            p.group_number,
            p.member_id,
            _street_inline(p),
            a["city"].title(),
            a["state"],
            a["zip5"],
        ]
        for d in deps.get(p.pid, [])[:3]:
            values += [d.first, d.last, d.dob, d.relationship.title()]
        for j, val in enumerate(values, start=1):
            c = ws.cell(row=r, column=j, value=val)
            if isinstance(val, date):
                c.number_format = DATE_FMT
        m = _mention(doc_file, sheet, r, "SUBSCRIBER", p, planted=tags | _unit_tag(p))
        m.fields = {
            **name_fields(p, with_middle=False),
            **a,
            "dob": p.dob.isoformat(),
            "ssn": p.ssn,
            "employee_id": p.employee_id,
            "group_number": p.group_number,
            "member_id": p.member_id,
        }
        m.printed = {"suffix": suffix}
        out.mentions.append(m)
        for d in deps.get(p.pid, [])[:3]:
            dt = {"P37"} | ({"P22"} if d.age_on(pop.reference_date) < 18 else set())
            dm = _mention(doc_file, sheet, r, "DEPENDENT", d, planted=dt, related=m)
            dm.fields = {
                **name_fields(d, with_middle=False),
                "dob": d.dob.isoformat(),
                "relationship": d.relationship,
            }
            out.mentions.append(dm)
    save_xlsx(wb, docs / doc_file, export)
    labels = [
        ("Employee ID", "EMPLOYEE_ID"),
        ("First Name", "FIRST_NAME"),
        ("Last Name", "LAST_NAME"),
        ("Suffix", "SUFFIX"),
        ("DOB", "DOB"),
        ("SSN", "SSN"),
        ("Plan", "OTHER"),
        ("Group No", "GROUP_NUMBER"),
        ("Member ID", "MEMBER_ID"),
        ("Address", "ADDRESS_LINE1"),
        ("City", "CITY"),
        ("State", "STATE"),
        ("ZIP", "ZIP"),
    ]
    for k in range(1, 4):
        labels += [
            (f"Dep{k}_First", "FIRST_NAME"),
            (f"Dep{k}_Last", "LAST_NAME"),
            (f"Dep{k}_DOB", "DOB"),
            (f"Dep{k}_Rel", "OTHER"),
        ]
    doc = DocRecord(
        doc_file,
        "XLSX",
        "SPREADSHEET",
        export,
        planted={"P37", "P22"},
        columns=_cols(sheet, labels),
    )
    doc.messes = [
        "wide dependent columns Dep1_First ... Dep3_Rel; one row holds a subscriber plus up to "
        "3 dependents (P37)",
        "minor dependents (P22); suffix spelled 'Junior' for one subscriber (P02)",
    ]
    out.docs.append(doc)
    return out


# ---- student_roster.xlsx ----------------------------------------------------------------
def _grade(p: Person, ref: date) -> str:
    age = p.age_on(ref)
    return "UG" if age >= 18 else ("K" if age <= 5 else str(min(12, age - 5)))


def render_students(pop: Population, s: SynthSettings, docs: Path, rng: random.Random) -> Rendered:
    name, sheet = "student_roster.xlsx", "Students"
    doc_file = f"sheets/{name}"
    export = s.config.sheets.export_dates[name]
    students = pop.by_cast(STUDENT) + [
        p for p in pop.people.values() if p.home_sheet == name and p.cast != EXTRA_GUARDIAN
    ]
    rng.shuffle(students)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    header = [
        "Student ID",
        "Student Name",
        "DOB",
        "Grade",
        "School",
        "Guardian Name",
        "Guardian Phone",
        "Address",
        "City",
        "State",
        "ZIP",
    ]
    _header(ws, 1, header)
    _notice(ws)
    out = Rendered()
    for i, st in enumerate(students):
        r = i + 2
        g = pop.people[st.guardian]
        a = address_fields(st.address)
        sn = name_fields(st)
        values: list[object] = [
            st.student_id,
            last_comma_first(sn).title(),
            st.dob,
            _grade(st, pop.reference_date),
            st.school.title(),
            f"{g.first} {g.last}",
            g.phone,
            _street_inline(st),
            a["city"].title(),
            a["state"],
            a["zip5"],
        ]
        for j, val in enumerate(values, start=1):
            c = ws.cell(row=r, column=j, value=val)
            if isinstance(val, date):
                c.number_format = DATE_FMT
        m = _mention(doc_file, sheet, r, "STUDENT", st, planted={"P40"} | _unit_tag(st))
        m.fields = {
            **sn,
            **a,
            "dob": st.dob.isoformat(),
            "student_id": st.student_id,
            "school": st.school,
        }
        m.printed = {"name": str(values[1])}
        gm = _mention(
            doc_file,
            sheet,
            r,
            "GUARDIAN",
            g,
            planted={"P23"} if g.cast == EXTRA_GUARDIAN else set(),
        )
        gm.fields = {**name_fields(g, with_middle=False), "phone": g.phone}
        gm.printed = {"name": f"{g.first} {g.last}"}
        m.related = gm
        out.mentions += [m, gm]
    save_xlsx(wb, docs / doc_file, export)
    labels = [
        ("Student ID", "STUDENT_ID"),
        ("Student Name", "FULL_NAME"),
        ("DOB", "DOB"),
        ("Grade", "OTHER"),
        ("School", "SCHOOL"),
        ("Guardian Name", "GUARDIAN_NAME"),
        ("Guardian Phone", "PHONE"),
        ("Address", "ADDRESS_LINE1"),
        ("City", "CITY"),
        ("State", "STATE"),
        ("ZIP", "ZIP"),
    ]
    doc = DocRecord(
        doc_file,
        "XLSX",
        "SPREADSHEET",
        export,
        planted={"P40", "P23"},
        columns=_cols(sheet, labels),
    )
    doc.messes = [
        "student records (student ID, school) -> FERPA flag; many students are minors (P40)",
        "guardian name + phone per row; most guardians appear nowhere else and have no "
        "qualifying element (P23)",
        "student name as 'Last, First M'",
    ]
    out.docs.append(doc)
    return out


# ---- customer_cards.xlsx ----------------------------------------------------------------
def render_customers(pop: Population, s: SynthSettings, docs: Path, rng: random.Random) -> Rendered:
    name, sheet = "customer_cards.xlsx", "Customers"
    doc_file = f"sheets/{name}"
    export = s.config.sheets.export_dates[name]
    bob = pop.role("nickname")
    cast = [p for p in pop.by_cast(BASE) if p.card and p.pid != bob.pid]
    people = [bob, *cast] + [p for p in pop.people.values() if p.home_sheet == name]
    rng.shuffle(people)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    _header(ws, 1, ["Customer", "Card Number", "Exp", "Street", "City/State/Zip", "Phone"])
    _notice(ws)
    out = Rendered()
    forced = ["sci", "premasked", "text"]
    for i, p in enumerate(people):
        r = i + 2
        tags: set[str] = set()
        kind = (
            forced[i]
            if i < len(forced)
            else rng.choices(["text", "sci", "premasked"], [0.6, 0.25, 0.15])[0]
        )
        if p.pid == pop.roles["luhn_fail"]:
            kind = "text"
            tags.add("P14")
        card_field = p.card
        cell_val: object
        if kind == "sci":  # P33: stored as a float, digits beyond 15 lost
            cell_val = float(int(p.card[:15]) * 10)
            card_field, tags = "", tags | {"P33"}
        elif kind == "premasked":  # P34
            cell_val = rng.choice([f"XXXX-XXXX-XXXX-{p.card[-4:]}", f"************{p.card[-4:]}"])
            card_field, tags = "", tags | {"P34"}
        else:
            cell_val = p.card
        first = s.config.planted.nickname.printed if p.pid == bob.pid else p.first
        if p.pid == bob.pid:
            tags.add("P05")
        a = address_fields(p.address)
        csz = f"{a['city']}, {a['state']} {a['zip5']}"
        exp = f"{rng.randint(1, 12):02d}/{rng.randint(27, 31)}"
        values = [f"{first} {p.last}", cell_val, exp, _street_inline(p), csz, p.phone]
        for j, val in enumerate(values, start=1):
            c = ws.cell(row=r, column=j, value=val)
            if j == 2 and isinstance(val, float):
                c.number_format = "0.00E+00"
        m = _mention(doc_file, sheet, r, "CUSTOMER", p, planted=tags | _unit_tag(p))
        m.fields = {
            **name_fields(p, first=first, with_middle=False),
            **a,
            "card": card_field,
            "card_status": {"sci": "PRECISION_LOST", "premasked": "PREMASKED"}.get(kind, "TEXT"),
            "phone": p.phone,
        }
        m.printed = {"name": str(values[0]), "card": str(cell_val), "city_state_zip": csz}
        out.mentions.append(m)
    save_xlsx(wb, docs / doc_file, export)
    labels = [
        ("Customer", "FULL_NAME"),
        ("Card Number", "CARD"),
        ("Exp", "OTHER"),
        ("Street", "ADDRESS_LINE1"),
        ("City/State/Zip", "CITY_STATE_ZIP"),
        ("Phone", "PHONE"),
    ]
    doc = DocRecord(
        doc_file,
        "XLSX",
        "SPREADSHEET",
        export,
        planted={"P33", "P34", "P14", "P05"},
        columns=_cols(sheet, labels),
    )
    doc.messes = [
        "card numbers stored as floats shown in scientific notation, last digit lost (P33)",
        "pre-masked cards 'XXXX-XXXX-XXXX-1234' / '************1234' (P34)",
        "one card fails the Luhn check (P14)",
        "'City, ST ZIP' combined in one column; nickname 'Bob' for Robert (P05)",
    ]
    out.docs.append(doc)
    return out


# ---- vendor_list.xlsx -------------------------------------------------------------------
def render_vendors(pop: Population, s: SynthSettings, docs: Path, rng: random.Random) -> Rendered:
    name, sheet = "vendor_list.xlsx", "Vendors"
    doc_file = f"sheets/{name}"
    export = s.config.sheets.export_dates[name]
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    _header(ws, 1, ["Vendor", "EIN", "Address", "City", "State", "ZIP", "Phone", "Category"])
    _notice(ws)
    out = Rendered()
    cats = ["Office supplies", "IT services", "Facilities", "Staffing", "Logistics", "Consulting"]
    orgs = sorted(pop.orgs.values(), key=lambda o: o.org_id)
    for i, o in enumerate(orgs[: s.config.n_vendor_rows]):
        r = i + 2
        a = address_fields(o.address)
        for j, val in enumerate(
            [
                o.name,
                o.ein,
                a["line1"],
                a["city"].title(),
                a["state"],
                a["zip5"],
                o.phone,
                rng.choice(cats),
            ],
            start=1,
        ):
            ws.cell(row=r, column=j, value=val)
        out.exclusions += [
            Exclusion(doc_file, "ORG", o.name, sheet=sheet, row=r, planted={"P39"}),
            Exclusion(doc_file, "EIN", o.ein, sheet=sheet, row=r, planted={"P39"}),
        ]
    save_xlsx(wb, docs / doc_file, export)
    labels = [
        ("Vendor", "ORG_NAME"),
        ("EIN", "EIN"),
        ("Address", "ADDRESS_LINE1"),
        ("City", "CITY"),
        ("State", "STATE"),
        ("ZIP", "ZIP"),
        ("Phone", "PHONE"),
        ("Category", "OTHER"),
    ]
    doc = DocRecord(
        doc_file, "XLSX", "SPREADSHEET", export, planted={"P39"}, columns=_cols(sheet, labels)
    )
    doc.messes = ["organizations only, with EINs: must produce 0 person mentions (P39)"]
    out.docs.append(doc)
    return out


# ---- legacy_export.xls ------------------------------------------------------------------
def soffice_path() -> str | None:
    env = os.environ.get("SOFFICE_PATH", "").strip()
    if env and Path(env).is_file():
        return env
    return shutil.which("soffice")


def convert_to_xls(src: Path, dst: Path) -> None:
    """xlsx -> Excel 97 .xls with LibreOffice (dedicated profile, ADR-041)."""
    exe = soffice_path()
    if exe is None:
        raise RuntimeError("LibreOffice not found: set SOFFICE_PATH or put soffice on PATH")
    with tempfile.TemporaryDirectory() as tmp:
        profile = Path(tmp, "lo_profile").as_uri()
        subprocess.run(
            [
                exe,
                f"-env:UserInstallation={profile}",
                "--headless",
                "--convert-to",
                "xls:MS Excel 97",
                "--outdir",
                tmp,
                str(src),
            ],
            check=True,
            capture_output=True,
            timeout=180,
        )
        produced = Path(tmp, src.with_suffix(".xls").name)
        dst.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(produced, dst)


def render_legacy(pop: Population, s: SynthSettings, docs: Path, rng: random.Random) -> Rendered:
    name, sheet = "legacy_export.xls", "Sheet1"
    doc_file = f"sheets/{name}"
    export = s.config.sheets.export_dates[name]
    variant = pop.role("variant_surname")
    pool = [p for p in pop.by_cast(BASE) if p.pid not in {pop.roles["mover"], variant.pid}]
    people = [variant, *rng.sample(pool, 24)] + [
        p for p in pop.people.values() if p.home_sheet == name
    ]
    rng.shuffle(people)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    _header(ws, 1, ["ID", "NAME", "ADDR1", "CITY", "ST", "ZIP", "Column7", "PHONE"])
    _notice(ws)
    out = Rendered()
    printed_last = up(s.config.planted.variant_surname.printed)
    for i, p in enumerate(people):
        r = i + 2
        tags = {"P36", "P04"}
        last = printed_last if p.pid == variant.pid else None
        if last:
            tags.add("P08")
        if p.pid == pop.roles.get("ssn_group_00"):
            tags.add("P13")
        n = name_fields(p, last=last)
        a = address_fields(p.address)
        values: list[object] = [
            p.ssn,
            last_comma_first(n),
            _street_inline(p),
            a["city"],
            a["state"],
            a["zip5"],
            p.dob,
            p.phone.removeprefix("+1"),
        ]
        for j, val in enumerate(values, start=1):
            c = ws.cell(row=r, column=j, value=val)
            if isinstance(val, date):
                c.number_format = DATE_FMT
        m = _mention(doc_file, sheet, r, "MEMBER", p, planted=tags | _unit_tag(p))
        m.fields = {**n, **a, "ssn": p.ssn, "dob": p.dob.isoformat(), "phone": p.phone}
        m.printed = {"name": str(values[1])}
        out.mentions.append(m)
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp, "legacy_export.xlsx")
        save_xlsx(wb, src, export)
        convert_to_xls(src, docs / doc_file)
    labels = [
        ("ID", "SSN"),
        ("NAME", "FULL_NAME"),
        ("ADDR1", "ADDRESS_LINE1"),
        ("CITY", "CITY"),
        ("ST", "STATE"),
        ("ZIP", "ZIP"),
        ("Column7", "DOB"),
        ("PHONE", "PHONE"),
    ]
    doc = DocRecord(
        doc_file,
        "XLS",
        "SPREADSHEET",
        export,
        planted={"P36", "P04", "P08", "P13"},
        columns=_cols(sheet, labels),
    )
    doc.messes = [
        "old .xls format (Excel 97, converted by LibreOffice)",
        "header 'ID' actually holds SSNs; 'Column7' holds DOBs (P36)",
        "NAME as 'LAST, FIRST M' (P04); one surname printed ANDERSEN for ANDERSON (P08)",
        "one SSN with group 00 (P13)",
    ]
    out.docs.append(doc)
    return out


SHEET_RENDERERS = [
    render_hr_roster,
    render_payroll,
    render_patients,
    render_benefits,
    render_students,
    render_customers,
    render_vendors,
    render_legacy,
]


def render_all_sheets(pop: Population, s: SynthSettings, docs: Path) -> Rendered:
    out = Rendered()
    for fn in SHEET_RENDERERS:
        out.extend(fn(pop, s, docs, random.Random(f"{s.config.seed}:{fn.__name__}")))
    return out


__all__ = ["SHEET_RENDERERS", "excel_serial", "first_last", "render_all_sheets", "soffice_path"]
