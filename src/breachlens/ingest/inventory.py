"""Inventory: one row per input file, both tracks (DATA_SCHEMA.md §inventory).

Input: the input directory. Output columns: doc_id, path, file_kind, sheets, sha256,
bytes, pages, has_text_layer, is_exact_duplicate, doc_type, doc_type_score,
extraction_method, mean_ocr_conf, processing_ms, doc_date. The last six are filled in by
later stages (PDF track: doc type, method, OCR confidence, timing, doc_date).
"""

from __future__ import annotations

import csv
import hashlib
import logging
from pathlib import Path
from typing import Any

import pandas as pd
import pymupdf
import xlrd
from openpyxl import load_workbook

from breachlens.ingest.doc_ids import assign_doc_ids, relative_files

logger = logging.getLogger(__name__)

KINDS = {".pdf": "PDF", ".xlsx": "XLSX", ".xls": "XLS", ".csv": "CSV"}
COLUMNS = [
    "doc_id",
    "path",
    "file_kind",
    "sheets",
    "sha256",
    "bytes",
    "pages",
    "has_text_layer",
    "is_exact_duplicate",
    "doc_type",
    "doc_type_score",
    "extraction_method",
    "mean_ocr_conf",
    "processing_ms",
    "doc_date",
]
MIN_TEXT_CHARS = 20


def pdf_facts(path: Path) -> tuple[int, bool]:
    """(pages, has_text_layer). A page with fewer than 20 text characters counts as image-only."""
    with pymupdf.open(path) as doc:
        chars = [len(page.get_text().strip()) for page in doc]
    return len(chars), any(n >= MIN_TEXT_CHARS for n in chars)


def sheet_count(path: Path, kind: str) -> int:
    if kind == "XLSX":
        wb = load_workbook(path, read_only=True)
        try:
            return len(wb.sheetnames)
        finally:
            wb.close()
    if kind == "XLS":
        return xlrd.open_workbook(path, on_demand=True).nsheets
    return 1


def build_inventory(input_dir: Path) -> pd.DataFrame:
    """Scan `input_dir` and return the inventory, sorted by doc_id."""
    rows: list[dict[str, Any]] = []
    seen: dict[str, str] = {}
    for rel, doc_id in sorted(
        assign_doc_ids(relative_files(input_dir)).items(), key=lambda kv: kv[1]
    ):
        path = input_dir / rel
        data = path.read_bytes()
        sha = hashlib.sha256(data).hexdigest()
        kind = KINDS.get(path.suffix.lower(), "OTHER")
        pages, text_layer, sheets = None, None, None
        if kind == "PDF":
            pages, text_layer = pdf_facts(path)
        elif kind in {"XLSX", "XLS", "CSV"}:
            try:
                sheets = sheet_count(path, kind)
            except (OSError, ValueError, xlrd.XLRDError, csv.Error) as exc:
                logger.warning("%s: cannot count sheets (%s)", doc_id, type(exc).__name__)
        rows.append(
            {
                "doc_id": doc_id,
                "path": rel,
                "file_kind": kind,
                "sheets": sheets,
                "sha256": sha,
                "bytes": len(data),
                "pages": pages,
                "has_text_layer": text_layer,
                "is_exact_duplicate": sha in seen,
                "doc_type": None,
                "doc_type_score": None,
                "extraction_method": None,
                "mean_ocr_conf": None,
                "processing_ms": None,
                "doc_date": None,
            }
        )
        seen.setdefault(sha, doc_id)
    df = pd.DataFrame(rows, columns=COLUMNS)
    logger.info(
        "inventory: %d files (%d PDF, %d duplicates)",
        len(df),
        int((df["file_kind"] == "PDF").sum()),
        int(df["is_exact_duplicate"].sum()),
    )
    return df
