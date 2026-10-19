"""Recalculate the review workbook in LibreOffice and check it (M7, ADR-041).

XlsxWriter stores formulas without results, so the workbook is converted headless with
a throwaway LibreOffice profile set to always recalculate on load. The result is read
with openpyxl (`data_only=True`) and checked:
- no `#REF!` / `#NAME?` / `#VALUE!` / `#DIV/0!` / `#N/A` / `#NUM!` / `#NULL!` anywhere
- every Summary formula equals the Python value from `summary_rows` (computed from the
  database, not read from the workbook)
- the Excel notify-status formula agrees with Python on every Notification_List row
"""

from __future__ import annotations

import os
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from breachlens.report.excel_writer import SUMMARY_HEADER_ROW, ReportData, summary_rows

ERROR_VALUES = {"#REF!", "#NAME?", "#VALUE!", "#DIV/0!", "#N/A", "#NUM!", "#NULL!"}

# OOXML/ODF "recalculation on file load": 0 = always (Tools > Options > Calc > Formula).
_LOAD = '<item oor:path="/org.openoffice.Office.Calc/Formula/Load">'
_XCU = (
    '<?xml version="1.0" encoding="UTF-8"?>\n'
    '<oor:items xmlns:oor="http://openoffice.org/2001/registry" '
    'xmlns:xs="http://www.w3.org/2001/XMLSchema" '
    'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">\n'
    + "".join(
        f'{_LOAD}<prop oor:name="{mode}" oor:op="fuse"><value>0</value></prop></item>\n'
        for mode in ("OOXMLRecalcMode", "ODFRecalcMode")
    )
    + "</oor:items>\n"
)


def soffice_path() -> str | None:
    env = os.environ.get("SOFFICE_PATH", "").strip()
    if env and Path(env).is_file():
        return env
    return shutil.which("soffice")


def recalculate(workbook: Path, out_dir: Path, timeout: int = 300) -> Path:
    """Convert `workbook` to a recalculated `.xlsx` in `out_dir` (dedicated profile)."""
    exe = soffice_path()
    if exe is None:
        raise RuntimeError("LibreOffice not found: set SOFFICE_PATH or put soffice on PATH")
    out_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        profile = Path(tmp, "lo_profile")
        (profile / "user").mkdir(parents=True)
        (profile / "user" / "registrymodifications.xcu").write_text(_XCU, encoding="utf-8")
        subprocess.run(
            [
                exe,
                f"-env:UserInstallation={profile.as_uri()}",
                "--headless",
                "--convert-to",
                "xlsx:Calc MS Excel 2007 XML",
                "--outdir",
                str(out_dir),
                str(workbook),
            ],
            check=True,
            capture_output=True,
            timeout=timeout,
        )
    produced = out_dir / f"{workbook.stem}.xlsx"
    if not produced.is_file():
        raise RuntimeError(f"LibreOffice produced no file for {workbook.name}")
    return produced


@dataclass
class RecalcResult:
    errors: list[tuple[str, str, str]] = field(default_factory=list)
    summary: list[tuple[str, str, object, object]] = field(default_factory=list)
    mismatches: list[tuple[str, str, object, object]] = field(default_factory=list)
    status_rows: int = 0
    status_disagree: int = 0

    @property
    def ok(self) -> bool:
        return not self.errors and not self.mismatches and self.status_disagree == 0


def _same(a: object, b: object) -> bool:
    if isinstance(a, bool) or isinstance(b, bool):
        return a == b
    try:
        return abs(float(a) - float(b)) < 1e-9  # type: ignore[arg-type]
    except (TypeError, ValueError):
        return False


def check(recalculated: Path, data: ReportData) -> RecalcResult:
    res = RecalcResult()
    wb = load_workbook(recalculated, data_only=True)
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.strip() in ERROR_VALUES:
                        res.errors.append((ws.title, c.coordinate, c.value))
        expected = {(r.section, r.item): r.python for r in summary_rows(data)}
        ws = wb["Summary"]
        seen = set()
        for row in ws.iter_rows(min_row=SUMMARY_HEADER_ROW + 2, values_only=True):
            section, item, excel = row[0], row[1], row[2]
            if (section, item) not in expected:
                continue
            seen.add((section, item))
            want = expected[(section, item)]
            res.summary.append((section, item, excel, want))
            if not _same(excel, want):
                res.mismatches.append((section, item, excel, want))
        for key in expected.keys() - seen:
            res.mismatches.append((key[0], key[1], None, expected[key]))
        ws = wb["Notification_List"]
        header = [c.value for c in ws[1]]
        col = header.index("status_agrees")
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] is None:
                continue
            res.status_rows += 1
            res.status_disagree += row[col] is not True
    finally:
        wb.close()
    return res
