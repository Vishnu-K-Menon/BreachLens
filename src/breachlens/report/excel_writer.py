"""Excel review workbook (M7, EXCEL_WORKBOOK.md as reduced by ADR-051/068).

Built from the DuckDB system of record (ADR-067) plus the run's guarded excluded_*
CSVs. Seven sheets: Summary, Mentions, Review_Queue, People, Notification_List,
Excluded, About. Every data block is an Excel Table, so the formulas use structured
references. Summary lists each count three ways: the live Excel formula, the Python
value, and whether they match; `report/recalc.py` recalculates the file in LibreOffice
and checks the formulas against `summary_rows` (the Python side).

Masked values only: `*_key` columns stay in the database, and every sheet is scanned by
the leak guard before it is written and again after the file is read back.
"""

from __future__ import annotations

import json
import math
import os
import re
import tempfile
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

import pandas as pd
import xlsxwriter
from openpyxl import load_workbook

from breachlens.db.sql_harness import connect_read_only
from breachlens.privacy.leak_guard import Finding, LeakError, LeakGuard

SHEETS = ("Summary", "Mentions", "Review_Queue", "People", "Notification_List", "Excluded", "About")
TABLES = {
    "Mentions": "tblMentions",
    "Review_Queue": "tblReview",
    "People": "tblPeople",
    "Notification_List": "tblNotify",
    "Excluded": "tblExcluded",
}
REVIEW_STATUSES = ["PENDING", "ACCEPTED", "REJECTED", "RESOLVED"]
REVIEW_FIRST_ROW = 3  # 0-based header row of tblReview; the macro buttons sit above it
TEXT_COLUMNS = re.compile(r"(_id|zip5|zip4|_masked|_valid|source_row|page|phone|item_id)$")

# Reason-code prefix -> the Mentions columns its fill colours (presentation only).
CODE_COLUMNS: list[tuple[tuple[str, ...], tuple[str, ...]]] = [
    (
        ("NAME_",),
        (
            "name_title",
            "first_name",
            "middle_name",
            "middle_initial",
            "last_name",
            "name_suffix",
        ),
    ),
    (
        ("ADDR_",),
        ("address_line1", "address_line2", "city", "state", "zip5", "zip4", "address_status"),
    ),
    (("DOB_",), ("dob", "age_at_run", "is_minor")),
    (("SSN_", "ITIN_", "XL_SSN_"), ("ssn_present", "ssn_masked", "ssn_valid")),
    (("DL_",), ("dl_present", "dl_masked", "dl_valid", "dl_state")),
    (("PASSPORT_",), ("passport_present", "passport_masked", "passport_valid")),
    (("CARD_", "XL_NUMBER_PRECISION"), ("card_present", "card_masked", "card_valid")),
    (
        ("ROUTING_",),
        (
            "bank_account_present",
            "bank_account_masked",
            "bank_account_valid",
            "routing_number_valid",
        ),
    ),
]
ADDRESS_COLUMNS = ("address_line1", "address_line2", "city", "state", "zip5", "zip4")


# Columns the formulas reference; a missing one is added empty so no formula hits #REF!.
REQUIRED = {
    "people": (
        "person_id",
        "first_name",
        "middle_initial",
        "last_name",
        "name_suffix",
        "state",
        "dob",
        "is_minor",
        "ssn_exposed",
        "medical_exposed",
        "elements_exposed",
        "ferpa_flag",
        "needs_review",
        "notify_required",
    ),
    "mentions": (
        "mention_id",
        "doc_id",
        "qc_status",
        "qc_reasons",
        "address_status",
        "review_status",
    ),
    "notification_list": (
        "person_id",
        "jurisdiction",
        "first_name",
        "middle_initial",
        "last_name",
        "name_suffix",
        "notify_required",
        "notify_status",
    ),
}


def _ensure(df: pd.DataFrame, cols: tuple[str, ...]) -> pd.DataFrame:
    missing = [c for c in cols if c not in df.columns]
    return df.assign(**{c: None for c in missing}) if missing else df


@dataclass
class ReportData:
    run_id: str
    as_of_date: date
    schema_version: str
    source_dir: str
    people: pd.DataFrame
    mentions: pd.DataFrame
    notify: pd.DataFrame
    review: pd.DataFrame
    excluded: pd.DataFrame


def load_report_data(db: Path, root: Path, as_of_date: date) -> ReportData:
    """Read the workbook's tables from the database (read-only) and the run folder."""
    if not db.is_file():
        raise FileNotFoundError(f"{db} not found; run `breachlens db build` first")
    con = connect_read_only(db)
    try:
        version, run_id, source_dir = con.execute(
            "SELECT schema_version, run_id, source_dir FROM schema_version"
        ).fetchone()
        tables = {
            t: _ensure(con.execute(f"SELECT * FROM {t}").df(), REQUIRED.get(t, ()))
            for t in ("people", "mentions", "notification_list", "review_items")
        }
    finally:
        con.close()
    run_dir = Path(source_dir) if Path(source_dir).is_absolute() else root / source_dir
    summary = run_dir / "run_summary.json"
    if summary.is_file():
        as_of_date = date.fromisoformat(json.loads(summary.read_text("utf-8"))["as_of_date"])
    mentions = tables["mentions"]
    mentions = mentions[[c for c in mentions.columns if not c.endswith("_key")]]
    return ReportData(
        run_id=run_id,
        as_of_date=as_of_date,
        schema_version=version,
        source_dir=source_dir,
        people=tables["people"].drop(columns="run_id"),
        mentions=mentions.drop(columns="run_id"),
        notify=tables["notification_list"].drop(columns="run_id"),
        review=review_queue(tables["review_items"]),
        excluded=excluded(run_dir),
    )


def review_queue(items: pd.DataFrame) -> pd.DataFrame:
    """ADR-062 order: OPEN first, ERROR before WARN, then most affected rows first."""
    q = items.assign(
        _open=(items["status"] != "OPEN").astype(int),
        _sev=(items["severity"] != "ERROR").astype(int),
    ).sort_values(["_open", "_sev", "affected_rows", "item_id"], ascending=[1, 1, 0, 1])
    q = q[
        [
            "item_type",
            "item_id",
            "severity",
            "codes",
            "affected_rows",
            "doc_id",
            "sheet",
            "header",
            "mention_id",
            "person_id",
            "status",
        ]
    ].reset_index(drop=True)
    q["review_status"] = q["status"].map(lambda s: "PENDING" if s == "OPEN" else "RESOLVED")
    for c in ("reviewer_note", "reviewed_by", "reviewed_at"):
        q[c] = None
    return q


def excluded(run_dir: Path) -> pd.DataFrame:
    frames = []
    for name in ("excluded_entities", "excluded_identifiers"):
        p = run_dir / f"{name}.csv"
        if p.is_file():
            df = pd.read_csv(p, dtype=str, keep_default_na=False)
            frames.append(df.assign(table=name.removeprefix("excluded_")))
    cols = ["table", "kind", "value_masked", "reason", "doc_id", "page", "sheet", "row"]
    if not frames:
        return pd.DataFrame(columns=cols)
    df = pd.concat(frames, ignore_index=True)
    df["row"] = df["row"].str.replace(r"\.0$", "", regex=True)
    return df[[c for c in cols if c in df.columns]]


# ---- Summary: one list drives both the formulas and the Python check ------------------
@dataclass(frozen=True)
class SummaryRow:
    section: str
    item: str
    formula: str  # `{row}` = this row's Excel row number, `{ref:<item>}` = another row's C cell
    python: float | int


def _b(s: pd.Series) -> pd.Series:
    return s.map(lambda v: False if v is None or v is pd.NA or v != v else bool(v))


def _age(dob: Any, as_of: date) -> int | None:
    if dob is None or (isinstance(dob, float) and math.isnan(dob)) or pd.isna(dob):
        return None
    d = pd.Timestamp(dob).date()
    return as_of.year - d.year - ((as_of.month, as_of.day) < (d.month, d.day))


def code_counts(mentions: pd.DataFrame) -> pd.Series:
    codes = mentions["qc_reasons"].fillna("").str.split(";").explode()
    return codes[codes != ""].value_counts()


def summary_rows(d: ReportData, top_codes: int = 10) -> list[SummaryRow]:
    p, n, r = d.people, d.notify, d.review
    ages = [_age(x, d.as_of_date) for x in p["dob"]]
    rows = [
        SummaryRow("KPI", "People", "=ROWS(tblPeople[person_id])", len(p)),
        SummaryRow("KPI", "Mentions", "=ROWS(tblMentions[mention_id])", len(d.mentions)),
        SummaryRow(
            "KPI",
            "Collapse ratio (mentions per person)",
            "=ROUND({ref:Mentions}/{ref:People},2)",
            round(len(d.mentions) / max(len(p), 1), 2),
        ),
        SummaryRow(
            "KPI",
            "People to notify",
            "=COUNTIFS(tblPeople[notify_required],TRUE)",
            int(_b(p["notify_required"]).sum()),
        ),
    ]
    for status in ("READY", "NEEDS_REVIEW", "NOT_REQUIRED"):
        rows.append(
            SummaryRow(
                "KPI",
                f"Notify status {status}",
                f'=COUNTIFS(tblNotify[notify_status],"{status}")',
                int((n["notify_status"] == status).sum()),
            )
        )
    rows += [
        SummaryRow(
            "KPI",
            "Excel status formula agrees with Python",
            "=COUNTIFS(tblNotify[status_agrees],TRUE)",
            len(n),
        ),
        SummaryRow(
            "KPI",
            "People needing review",
            "=COUNTIFS(tblPeople[needs_review],TRUE)",
            int(_b(p["needs_review"]).sum()),
        ),
        SummaryRow(
            "KPI",
            "Minors (flag)",
            "=COUNTIFS(tblPeople[is_minor],TRUE)",
            int(_b(p["is_minor"]).sum()),
        ),
        SummaryRow(
            "KPI",
            "Minors (DATEDIF from RunDate)",
            '=COUNTIFS(tblPeople[age_at_rundate],"<18")',
            sum(1 for a in ages if a is not None and a < 18),
        ),
        SummaryRow(
            "KPI",
            "FERPA flagged",
            "=COUNTIFS(tblPeople[ferpa_flag],TRUE)",
            int(_b(p["ferpa_flag"]).sum()),
        ),
        SummaryRow(
            "KPI",
            "SSN + medical exposed",
            "=SUMPRODUCT((tblPeople[ssn_exposed]=TRUE)*(tblPeople[medical_exposed]=TRUE))",
            int((_b(p["ssn_exposed"]) & _b(p["medical_exposed"])).sum()),
        ),
        SummaryRow(
            "KPI",
            "Review items open",
            '=COUNTIFS(tblReview[status],"OPEN")',
            int((r["status"] == "OPEN").sum()),
        ),
    ]
    for t in ("COLUMN", "MENTION", "PAIR"):
        rows.append(
            SummaryRow(
                "KPI",
                f"Open {t} items",
                f'=COUNTIFS(tblReview[item_type],"{t}",tblReview[status],"OPEN")',
                int(((r["item_type"] == t) & (r["status"] == "OPEN")).sum()),
            )
        )
    col_open = (r["item_type"] == "COLUMN") & (r["status"] == "OPEN")
    rows += [
        SummaryRow(
            "KPI",
            "Rows affected by open COLUMN items",
            '=SUMIFS(tblReview[affected_rows],tblReview[item_type],"COLUMN",tblReview[status],"OPEN")',
            int(r.loc[col_open, "affected_rows"].sum()),
        ),
        SummaryRow(
            "KPI",
            "Resolved COLUMN items (review_overrides.yaml)",
            '=COUNTIFS(tblReview[item_type],"COLUMN",tblReview[status],"RESOLVED")',
            int(((r["item_type"] == "COLUMN") & (r["status"] == "RESOLVED")).sum()),
        ),
        SummaryRow(
            "KPI",
            "Review items PENDING (reviewer column)",
            '=COUNTIFS(tblReview[review_status],"PENDING")',
            int((r["review_status"] == "PENDING").sum()),
        ),
        SummaryRow(
            "KPI", "Excluded entities/identifiers", "=COUNTA(tblExcluded[kind])", len(d.excluded)
        ),
    ]
    for st in sorted(n["jurisdiction"].dropna().unique()):
        rows.append(
            SummaryRow(
                "To notify by state",
                st,
                "=COUNTIFS(tblNotify[jurisdiction],$B{row},tblNotify[notify_required],TRUE)",
                int(((n["jurisdiction"] == st) & _b(n["notify_required"])).sum()),
            )
        )
    for st in sorted(p["state"].dropna().unique()):
        rows.append(
            SummaryRow(
                "SSN + medical by state",
                st,
                "=SUMPRODUCT((tblPeople[state]=$B{row})*(tblPeople[ssn_exposed]=TRUE)"
                "*(tblPeople[medical_exposed]=TRUE))",
                int(((p["state"] == st) & _b(p["ssn_exposed"]) & _b(p["medical_exposed"])).sum()),
            )
        )
    counts = code_counts(d.mentions)
    top = sorted(counts.items(), key=lambda kv: (-kv[1], kv[0]))[:top_codes]
    for code, k in top:
        rows.append(
            SummaryRow(
                "Top reason codes (mentions)",
                code,
                '=IFERROR(SUMPRODUCT(--ISNUMBER(SEARCH(";"&$B{row}&";",";"&tblMentions[qc_reasons]&";"))),0)',
                int(k),
            )
        )
    return rows


SUMMARY_HEADER_ROW = 5  # 0-based; data rows follow


# ---- writing ------------------------------------------------------------------------
def _cell(v: Any) -> Any:
    if v is None or v is pd.NA or v is pd.NaT:
        return None
    if isinstance(v, float) and math.isnan(v):
        return None
    if hasattr(v, "item") and not isinstance(v, (str, bytes)):
        v = v.item()
    if isinstance(v, pd.Timestamp):
        v = v.to_pydatetime()
    return v


def _rows(df: pd.DataFrame, date_cols: set[str]) -> list[list[Any]]:
    out = []
    for rec in df.itertuples(index=False, name=None):
        row = []
        for c, v in zip(df.columns, rec, strict=True):
            v = _cell(v)
            if c in date_cols and isinstance(v, datetime):
                v = v.date()
            row.append(v)
        out.append(row)
    return out


class _Formats:
    def __init__(self, wb: xlsxwriter.Workbook) -> None:
        self.text = wb.add_format({"num_format": "@"})
        self.date = wb.add_format({"num_format": "yyyy-mm-dd"})
        self.ts = wb.add_format({"num_format": "yyyy-mm-dd hh:mm:ss"})
        self.ratio = wb.add_format({"num_format": "0.00"})
        self.bold = wb.add_format({"bold": True})
        self.title = wb.add_format({"bold": True, "font_size": 14})
        self.section = wb.add_format({"bold": True, "bg_color": "#DDEBF7"})
        self.red = wb.add_format({"bg_color": "#FFC7CE", "font_color": "#9C0006"})
        self.yellow = wb.add_format({"bg_color": "#FFEB9C", "font_color": "#9C5700"})
        self.green = wb.add_format({"bg_color": "#C6EFCE", "font_color": "#006100"})
        self.grey = wb.add_format({"bg_color": "#D9D9D9", "font_color": "#595959"})
        self.wrap = wb.add_format({"text_wrap": True, "valign": "top"})


def _add_table(
    ws: Any,
    name: str,
    df: pd.DataFrame,
    f: _Formats,
    first_row: int = 0,
    formulas: dict[str, str] | None = None,
) -> tuple[int, int, dict[str, int]]:
    """Write `df` as an Excel Table; returns (first data row, last row, column index)."""
    date_cols = {c for c in df.columns if pd.api.types.is_datetime64_any_dtype(df[c])}
    date_only = {c for c in date_cols if c in {"dob", "doc_date"}}
    columns = []
    for c in df.columns:
        spec: dict[str, Any] = {"header": c}
        if c in date_only:
            spec["format"] = f.date
        elif c in date_cols:
            spec["format"] = f.ts
        elif TEXT_COLUMNS.search(c) and not pd.api.types.is_bool_dtype(df[c]):
            spec["format"] = f.text
        columns.append(spec)
    for c, formula in (formulas or {}).items():
        # `[@col]` -> `tblX[@[col]]`: the qualified form LibreOffice also resolves
        columns.append({"header": c, "formula": re.sub(r"\[@(\w+)\]", rf"{name}[@[\1]]", formula)})
    data = _rows(df, date_only)
    last_row = first_row + max(len(data), 1)
    options: dict[str, Any] = {"name": name, "columns": columns, "style": "Table Style Light 9"}
    if data:
        options["data"] = data
    ws.add_table(first_row, 0, last_row, len(columns) - 1, options)
    ws.freeze_panes(first_row + 1, 1)
    for i, spec in enumerate(columns):
        ws.set_column(i, i, min(max(len(spec["header"]) + 2, 10), 28))
    return first_row + 1, last_row, {spec["header"]: i for i, spec in enumerate(columns)}


def _col(idx: int) -> str:
    return xlsxwriter.utility.xl_col_to_name(idx)


def _write_mentions(ws: Any, d: ReportData, f: _Formats, severity: dict[str, str]) -> None:
    r1, r2, cols = _add_table(ws, TABLES["Mentions"], d.mentions, f)
    qc = f"${_col(cols['qc_reasons'])}{r1 + 1}"
    for prefixes, targets in CODE_COLUMNS:
        idx = [cols[c] for c in targets if c in cols]
        if not idx:
            continue
        for sev, fmt in (("ERROR", f.red), ("WARN", f.yellow)):
            codes = sorted(c for c, s in severity.items() if s == sev and c.startswith(prefixes))
            if not codes:
                continue
            tests = ",".join(f'ISNUMBER(SEARCH(";{c};",";"&{qc}&";"))' for c in codes)
            for i in idx:
                ws.conditional_format(
                    r1, i, r2, i, {"type": "formula", "criteria": f"=OR({tests})", "format": fmt}
                )
    status = f"${_col(cols['address_status'])}{r1 + 1}"
    for c in ADDRESS_COLUMNS:
        if c in cols:
            ws.conditional_format(
                r1,
                cols[c],
                r2,
                cols[c],
                {"type": "formula", "criteria": f'={status}="VERIFIED"', "format": f.green},
            )
    qs = f"${_col(cols['qc_status'])}{r1 + 1}"
    for value, fmt in (("FAIL", f.red), ("REVIEW", f.yellow)):
        ws.conditional_format(
            r1,
            cols["qc_status"],
            r2,
            cols["qc_status"],
            {"type": "formula", "criteria": f'={qs}="{value}"', "format": fmt},
        )
    rs = f"${_col(cols['review_status'])}{r1 + 1}"
    ws.conditional_format(
        r1,
        0,
        r2,
        len(cols) - 1,
        {"type": "formula", "criteria": f'={rs}="REJECTED"', "format": f.grey},
    )


def _write_review(ws: Any, d: ReportData, f: _Formats, buttons: bool) -> None:
    ws.write(0, 0, "Review queue: one row per column-level item (ADR-062)", f.title)
    ws.write(
        1,
        0,
        "Set review_status from the dropdown. RESOLVED = decided in config/review_overrides.yaml.",
    )
    go = (
        '=IFERROR(IF([@mention_id]<>"",'
        'HYPERLINK("#Mentions!A"&(MATCH([@mention_id],tblMentions[mention_id],0)+1),"go"),'
        'HYPERLINK("#Mentions!A"&(MATCH([@doc_id],tblMentions[doc_id],0)+1),"go")),"")'
    )
    r1, r2, cols = _add_table(
        ws, TABLES["Review_Queue"], d.review, f, first_row=REVIEW_FIRST_ROW, formulas={"go": go}
    )
    c = cols["review_status"]
    ws.data_validation(
        r1,
        c,
        r2,
        c,
        {
            "validate": "list",
            "source": REVIEW_STATUSES,
            "error_title": "review_status",
            "error_message": "Pick PENDING, ACCEPTED, REJECTED or RESOLVED.",
        },
    )
    ws.set_column(cols["codes"], cols["codes"], 40)
    ws.set_column(cols["item_id"], cols["item_id"], 34)
    ref = f"${_col(c)}{r1 + 1}"
    ws.conditional_format(
        r1,
        0,
        r2,
        len(cols) - 1,
        {"type": "formula", "criteria": f'={ref}="RESOLVED"', "format": f.grey},
    )
    sev = f"${_col(cols['severity'])}{r1 + 1}"
    ws.conditional_format(
        r1,
        cols["severity"],
        r2,
        cols["severity"],
        {"type": "formula", "criteria": f'={sev}="ERROR"', "format": f.red},
    )
    if buttons:
        ws.insert_button(
            "H1",
            {"macro": "ShowPendingOnly", "caption": "Show pending", "width": 120, "height": 28},
        )
        ws.insert_button(
            "J1",
            {"macro": "ExportDecisions", "caption": "Export decisions", "width": 130, "height": 28},
        )


def _write_summary(ws: Any, d: ReportData, f: _Formats, rows: list[SummaryRow]) -> None:
    ws.write(0, 0, "BreachLens review workbook - SYNTHETIC DATA, NOT REAL", f.title)
    ws.write(1, 0, "RunId", f.bold)
    ws.write_string(1, 1, d.run_id)
    ws.write(2, 0, "RunDate (as_of_date)", f.bold)
    ws.write_datetime(2, 1, datetime.combine(d.as_of_date, datetime.min.time()), f.date)
    ws.write(3, 0, "Excel = live formula; Python = value from the pipeline; match = Excel = Python")
    for i, h in enumerate(("section", "item", "excel", "python", "match")):
        ws.write(SUMMARY_HEADER_ROW, i, h, f.section)
    where = {}
    for k, r in enumerate(rows):
        where[r.item] = SUMMARY_HEADER_ROW + 1 + k
    for k, r in enumerate(rows):
        row = SUMMARY_HEADER_ROW + 1 + k
        formula = r.formula.replace("{row}", str(row + 1))
        formula = re.sub(r"\{ref:([^}]+)\}", lambda m: f"C{where[m.group(1)] + 1}", formula)
        ws.write_string(row, 0, r.section)
        ws.write_string(row, 1, r.item)
        ws.write_formula(row, 2, formula, f.ratio if isinstance(r.python, float) else None)
        ws.write_number(row, 3, r.python, f.ratio if isinstance(r.python, float) else None)
        ws.write_formula(row, 4, f"=C{row + 1}=D{row + 1}")
    last = SUMMARY_HEADER_ROW + len(rows)
    ws.conditional_format(
        SUMMARY_HEADER_ROW + 1,
        4,
        last,
        4,
        {"type": "cell", "criteria": "==", "value": "FALSE", "format": f.red},
    )
    ws.set_column(0, 0, 28)
    ws.set_column(1, 1, 44)
    ws.set_column(2, 4, 12)


ABOUT = [
    "SYNTHETIC DATA - NOT REAL. Every name, identifier and address was generated by "
    "breachlens.synth; any match to a real person or number is coincidental.",
    "NOT LEGAL ADVICE. Notification rules are an illustrative subset (config/jurisdictions.yaml).",
    "Masked values only: identifiers appear as *_present / *_masked / *_valid. Join keys "
    "(*_key) stay in the DuckDB database.",
    "",
    "Sheets: Summary (live formulas vs Python), Mentions, Review_Queue, People, "
    "Notification_List, Excluded, About.",
    "Review_Queue: column-level warnings are one item per column with affected_rows "
    "(ADR-062); columns decided in config/review_overrides.yaml show as RESOLVED (ADR-061).",
    "",
    "Macros (.xlsm only): 'Show pending' filters review_status = PENDING (click again to "
    "clear); 'Export decisions' writes the non-PENDING rows to review_decisions_<RunId>.csv "
    "next to this workbook.",
    "Enabling macros: a file downloaded from the internet opens with macros blocked (Mark of "
    "the Web). Close it, right-click the file > Properties > tick Unblock > OK, reopen it and "
    "click Enable Content.",
    "Formulas are live: Excel recalculates them on open. RunDate is the fixed as_of_date "
    "(ADR-053), never TODAY().",
    "",
    "Decisions: ADR-051 (Lite scope), ADR-062, ADR-067 (database), ADR-068 (this workbook).",
]


def _write_about(ws: Any, d: ReportData, f: _Formats) -> None:
    ws.write(0, 0, "About this workbook", f.title)
    meta = [
        ("Run id", d.run_id),
        ("Database schema version", d.schema_version),
        ("as_of_date (RunDate)", d.as_of_date.isoformat()),
        ("Source run folder", d.source_dir.replace("\\", "/")),
    ]
    for i, (k, v) in enumerate(meta, start=2):
        ws.write(i, 0, k, f.bold)
        ws.write_string(i, 1, v)
    for i, line in enumerate(ABOUT, start=len(meta) + 3):
        ws.write_string(i, 0, line)
    ws.set_column(0, 0, 26)
    ws.set_column(1, 1, 60)


def write_workbook(
    d: ReportData,
    out_dir: Path,
    guard: LeakGuard,
    severity: dict[str, str],
    vba: Path | None = None,
) -> Path:
    """Write the workbook (`.xlsm` with buttons if `vba` exists, else `.xlsx`).

    The sheets are scanned before writing; the file is built in a temp file, read back and
    scanned again, and only then moved into place. Any finding raises `LeakError`.
    """
    frames = {
        "Mentions": d.mentions,
        "Review_Queue": d.review,
        "People": d.people,
        "Notification_List": d.notify,
        "Excluded": d.excluded,
    }
    findings: list[Finding] = []
    for sheet, df in frames.items():
        findings += guard.scan_frame(df, sheet)
    if findings:
        raise LeakError(findings)
    macros = vba is not None and vba.is_file()
    out_dir.mkdir(parents=True, exist_ok=True)
    final = out_dir / f"breachlens_{d.run_id}.{'xlsm' if macros else 'xlsx'}"
    fd, tmp_name = tempfile.mkstemp(suffix=final.suffix, dir=out_dir, prefix=".tmp_")
    os.close(fd)
    tmp = Path(tmp_name)
    try:
        wb = xlsxwriter.Workbook(
            str(tmp),
            {
                "use_future_functions": True,
                "strings_to_formulas": False,
                "strings_to_urls": False,
                "remove_timezone": True,
            },
        )
        f = _Formats(wb)
        ws = {name: wb.add_worksheet(name) for name in SHEETS}
        rows = summary_rows(d)
        _write_summary(ws["Summary"], d, f, rows)
        wb.define_name("RunId", "=Summary!$B$2")
        wb.define_name("RunDate", "=Summary!$B$3")
        _write_mentions(ws["Mentions"], d, f, severity)
        _write_review(ws["Review_Queue"], d, f, buttons=macros)
        _add_table(
            ws["People"],
            TABLES["People"],
            d.people,
            f,
            formulas={
                "age_at_rundate": '=IF(ISNUMBER([@dob]),DATEDIF([@dob],RunDate,"y"),"")',
                "open_review_items": (
                    '=COUNTIFS(tblReview[person_id],[@person_id],tblReview[status],"OPEN")'
                ),
            },
        )
        _add_table(
            ws["Notification_List"],
            TABLES["Notification_List"],
            d.notify,
            f,
            formulas={
                "elements_xlookup": (
                    "=XLOOKUP([@person_id],tblPeople[person_id],"
                    'tblPeople[elements_exposed],"NOT FOUND")&""'
                ),
                "elements_index_match": (
                    '=INDEX(tblPeople[elements_exposed],MATCH([@person_id],tblPeople[person_id],0))&""'
                ),
                "mailing_name": (
                    '=TEXTJOIN(" ",TRUE,[@first_name],[@middle_initial],'
                    "[@last_name],[@name_suffix])"
                ),
                "status_xl": (
                    "=IF(AND([@notify_required],NOT(XLOOKUP([@person_id],tblPeople[person_id],"
                    'tblPeople[needs_review],FALSE))),"READY",'
                    'IF([@notify_required],"NEEDS_REVIEW","NOT_REQUIRED"))'
                ),
                "status_agrees": "=[@status_xl]=[@notify_status]",
            },
        )
        _add_table(ws["Excluded"], TABLES["Excluded"], d.excluded, f)
        _write_about(ws["About"], d, f)
        ws["Summary"].activate()
        if macros:
            wb.add_vba_project(str(vba))
        wb.close()
        if findings := scan_workbook(tmp, guard):
            raise LeakError(findings)
        os.replace(tmp, final)
    finally:
        tmp.unlink(missing_ok=True)
    return final


def scan_workbook(path: Path, guard: LeakGuard) -> list[Finding]:
    """Leak-scan every cell of a written workbook (values and formulas).

    Cells inside an Excel Table are scanned under their table header, so the guard's
    column rules (exemptions, account columns) apply; other cells are scanned as text.
    """
    findings: list[Finding] = []
    with path.open("rb") as fh:  # the handle closes even if openpyxl fails
        wb = load_workbook(fh, read_only=False, keep_vba=False)
    try:
        for ws in wb.worksheets:
            inside: set[tuple[int, int]] = set()
            for tbl in ws.tables.values():
                cells = ws[tbl.ref]
                header = [str(c.value) for c in cells[0]]
                body = [[c.value for c in row] for row in cells[1:]]
                for row in cells:
                    inside.update((c.row, c.column) for c in row)
                df = pd.DataFrame(body, columns=header, dtype=object)
                findings += guard.scan_frame(df, f"{path.name}:{ws.title}")
            for row in ws.iter_rows():
                for c in row:
                    if (c.row, c.column) not in inside and isinstance(c.value, str):
                        findings += guard.scan_text(
                            c.value, f"{path.name}:{ws.title}", c.coordinate
                        )
    finally:
        wb.close()
    return findings
