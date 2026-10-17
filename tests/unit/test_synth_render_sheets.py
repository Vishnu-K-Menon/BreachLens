import random
import re
from datetime import date, datetime
from pathlib import Path

import pytest
import xlrd
from openpyxl import load_workbook

from breachlens.synth import render_sheets as R
from breachlens.synth.people import build_population, read_pool
from breachlens.synth.settings import load_synth_settings

ROOT = Path(__file__).resolve().parents[2]


@pytest.fixture(scope="module")
def world():
    s = load_synth_settings(ROOT / "config")
    return s, build_population(s, read_pool(ROOT / "data" / "reference" / "address_pool.csv"))


def _render(fn, world, tmp_path):
    s, pop = world
    return fn(pop, s, tmp_path, random.Random(f"{s.config.seed}:{fn.__name__}"))


def _persons(out):
    return [m for m in out.mentions if m.person is not None]


def test_hr_roster_messes(world, tmp_path):
    out = _render(R.render_hr_roster, world, tmp_path)
    wb = load_workbook(tmp_path / "sheets" / "hr_roster.xlsx")
    ws = wb.active
    assert (
        ws["A3"].value.startswith("SYNTHETIC - NOT REAL")
        and ws["A1"].comment.text == "SYNTHETIC - NOT REAL"
    )
    merged = {str(r) for r in ws.merged_cells.ranges}
    assert {"A4:G4", "H4:K4", "N4:N5"} <= merged
    assert sum(1 for d in ws.row_dimensions.values() if d.hidden) == 3
    zips = [ws.cell(row=m.row, column=11).value for m in _persons(out)]
    assert all(isinstance(z, int) for z in zips) and any(z < 10000 for z in zips)
    assert any(
        str(ws.cell(row=r, column=1).value).startswith("Total")
        for r in range(ws.max_row - 1, ws.max_row + 1)
    )
    dup = next(m for m in out.mentions if "duplicate_of_row" in m.fields)
    src = int(dup.fields["duplicate_of_row"])
    assert [c.value for c in ws[dup.row]] == [c.value for c in ws[src]]
    assert any(m.entity_type == "ORG" for m in out.mentions)
    assert wb.properties.created == datetime(2026, 9, 30, 9, 0)


def test_save_xlsx_is_byte_identical(world, tmp_path):
    _render(R.render_vendors, world, tmp_path / "a")
    _render(R.render_vendors, world, tmp_path / "b")
    a = (tmp_path / "a" / "sheets" / "vendor_list.xlsx").read_bytes()
    assert a == (tmp_path / "b" / "sheets" / "vendor_list.xlsx").read_bytes()


def test_payroll_csv_messes(world, tmp_path):
    out = _render(R.render_payroll, world, tmp_path)
    raw = (tmp_path / "sheets" / "payroll_export.csv").read_bytes()
    with pytest.raises(UnicodeDecodeError):
        raw.decode("utf-8")
    lines = raw.decode("latin-1").split("\r\n")
    assert lines[0] == "# SYNTHETIC - NOT REAL" and "SS#" in lines[1].split(";")
    ssns = [line.split(";")[2] for line in lines[2:] if line]
    assert any(len(x) == 8 for x in ssns) and "#N/A" in raw.decode("latin-1")
    assert any("JJ," in line for line in lines)
    assert {m.row for m in out.mentions} == set(range(3, 3 + len(ssns)))


def test_patient_list_messes(world, tmp_path):
    out = _render(R.render_patients, world, tmp_path)
    wb = load_workbook(tmp_path / "sheets" / "patient_list.xlsx")
    assert wb.sheetnames == ["Patients", "Notes"]
    dobs = [wb["Patients"].cell(row=m.row, column=3).value for m in out.mentions]
    assert {type(v) for v in dobs} >= {int, str, datetime}
    notes = [wb["Patients"].cell(row=m.row, column=12).value or "" for m in out.mentions]
    assert any(re.search(r"\d{3}-\d{2}-\d{4}", n) for n in notes)
    assert any("MRN" in n for n in notes)


def test_benefits_unpivot_links_dependents(world, tmp_path):
    out = _render(R.render_benefits, world, tmp_path)
    deps = [m for m in out.mentions if m.role == "DEPENDENT"]
    assert len(deps) == 10 and all(d.related is not None and d.related.row == d.row for d in deps)
    assert any("P22" in d.planted for d in deps)
    ws = load_workbook(tmp_path / "sheets" / "benefits_enrollment.xlsx").active
    assert [c.value for c in ws[1]][13:17] == ["Dep1_First", "Dep1_Last", "Dep1_DOB", "Dep1_Rel"]


def test_student_roster_has_guardian_only_people(world, tmp_path):
    out = _render(R.render_students, world, tmp_path)
    guardians = [m for m in out.mentions if m.role == "GUARDIAN"]
    assert any("P23" in g.planted for g in guardians)
    assert all(
        set(g.fields) <= {"title", "first", "middle_initial", "last", "suffix", "phone"}
        for g in guardians
    )


def test_customer_cards_messes(world, tmp_path):
    out = _render(R.render_customers, world, tmp_path)
    ws = load_workbook(tmp_path / "sheets" / "customer_cards.xlsx").active
    cells = [ws.cell(row=m.row, column=2) for m in out.mentions]
    assert any(isinstance(c.value, (int, float)) and c.number_format == "0.00E+00" for c in cells)
    assert any(isinstance(c.value, str) and c.value.startswith(("XXXX", "****")) for c in cells)
    assert any("P14" in m.planted for m in out.mentions) and any(
        "P05" in m.planted for m in out.mentions
    )


def test_vendor_list_yields_no_person_mentions(world, tmp_path):
    out = _render(R.render_vendors, world, tmp_path)
    assert out.mentions == [] and {e.kind for e in out.exclusions} == {"ORG", "EIN"}


@pytest.mark.libreoffice
def test_legacy_export_is_real_xls(world, tmp_path):
    out = _render(R.render_legacy, world, tmp_path)
    path = tmp_path / "sheets" / "legacy_export.xls"
    assert path.read_bytes()[:8] == bytes.fromhex("D0CF11E0A1B11AE1")
    sh = xlrd.open_workbook(path).sheet_by_index(0)
    assert sh.row_values(0)[:7] == ["ID", "NAME", "ADDR1", "CITY", "ST", "ZIP", "Column7"]
    assert any("P08" in m.planted for m in out.mentions)
    assert R.excel_serial(date(1990, 1, 1)) == 32874
