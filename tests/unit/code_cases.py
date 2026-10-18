"""Reason-code case registry: every M4 code in config/rules.yaml gets >= 1 positive and
>= 1 negative case (VALIDATION_RULES.md). Each case calls the real producer (normalizer,
check, M3 spreadsheet producer) and returns the set of codes it fired.

`test_reason_codes.py` runs them; `counts()` gives the per-code table for the report.
"""

from __future__ import annotations

import json
from collections import Counter
from collections.abc import Callable
from dataclasses import dataclass
from datetime import date
from functools import cache
from pathlib import Path
from types import SimpleNamespace

import pandas as pd
from openpyxl import Workbook

from breachlens.address.zipcheck import check_zip
from breachlens.extract.entity_type import ORG, classify_entity, load_entity_config
from breachlens.normalize.addresses import normalize_address
from breachlens.normalize.names import normalize_parts, parse_full_name
from breachlens.sheets import repair as R
from breachlens.sheets.layout import fix_layout
from breachlens.sheets.mapping import decide, unmapped_sensitive
from breachlens.sheets.reader import read_file
from breachlens.sheets.settings import load_sheet_settings
from breachlens.sheets.sheet_track import map_columns, process_sheet
from breachlens.validate import completeness as C
from breachlens.validate.aba import routing_codes
from breachlens.validate.dl import dl_codes, passport_codes
from breachlens.validate.dob import dob_codes
from breachlens.validate.engine import doc_codes, exclusion_reason, name_variant_codes
from breachlens.validate.luhn import card_codes
from breachlens.validate.settings import load_validation_settings
from breachlens.validate.ssn import ssn_codes

ROOT = Path(__file__).resolve().parents[2]
V = load_validation_settings(ROOT / "config")
S = load_sheet_settings(ROOT / "config")
AS_OF = date(2026, 10, 1)
LEX_FIRST = frozenset({"LEE", "ALLEN", "JOHN", "MARY", "KELLY", "DEAN"})


@dataclass(frozen=True)
class Case:
    code: str
    positive: bool
    fn: Callable[[], set[str]]
    note: str = ""


def name(raw: str, order: str | None = None, labelled: bool = True) -> set[str]:
    return set(parse_full_name(raw, order, V.rules, LEX_FIRST, labelled).codes)


def parts(first: str, middle: str, last: str, suffix: str | None = None) -> set[str]:
    return set(normalize_parts(first, middle, last, suffix, V.rules).codes)


def addr(**row: str) -> set[str]:
    a = normalize_address(row, V.address, V.rules.zip_leading_zero_states)
    z = check_zip(a.zip5, a.city, a.state, V.rules.zip_city_suggest_min_similarity)
    return set(a.codes) | set(z.codes)


def dob(iso: str | None, raw: str | None = None) -> set[str]:
    return set(dob_codes(iso, raw, AS_OF, V.rules.max_age_years, V.rules.minor_age_years)[0])


def ssn(v: str) -> set[str]:
    return set(ssn_codes(v, V.identifiers, V.rules.known_invalid_ssns))


def variants(names: list[tuple[str, str]]) -> set[str]:
    df = pd.DataFrame(
        [
            {"doc_id": "D1", "_raw_ssn": "123456780", "first_name": f, "last_name": la}
            for f, la in names
        ]
    )
    return {"NAME_VARIANT_IN_DOC"} if name_variant_codes(df) else set()


def doc_type(doc: str) -> set[str]:
    inv = pd.DataFrame(
        [{"doc_id": "D1", "file_kind": "PDF", "doc_type": doc, "is_exact_duplicate": False}]
    )
    return set(doc_codes(SimpleNamespace(tables={"inventory": inv}), V)["code"])


def excluded(kind: str) -> set[str]:
    return {exclusion_reason(kind, "label: x").split(" ")[0]}


@cache
def _tmp() -> Path:
    import tempfile

    return Path(tempfile.mkdtemp(prefix="codecases_"))


def sheet(csv_text: str, encoding: str = "utf-8") -> set[str]:
    """Run a CSV through reader -> layout -> rules mapping -> row producers; all codes."""
    p = _tmp() / f"s{abs(hash((csv_text, encoding)))}.csv"
    p.write_bytes(csv_text.encode(encoding))
    [raw] = read_file(p, "D0001", p.name, "CSV", S.sheets.reader)
    t = fix_layout(raw, S.sheets.layout)
    ctx = SimpleNamespace(as_of_date=AS_OF, run_id="R20261001-000000-ab12")
    m = map_columns(t, S, None, ctx)
    out = process_sheet(t, m, S, ctx, load_entity_config(ROOT / "config" / "entities.yaml"))
    codes = {c for c, _ in t.codes} | {c for r in m.rows for c in r["codes"].split(";") if c}
    return codes | {c for r in out.mentions for c in r["_codes"]}


def xlsx_layout(
    hidden: bool = False,
    title: bool = False,
    merged: bool = False,
    total: bool = False,
    data: bool = True,
) -> set[str]:
    p = _tmp() / f"x{hidden}{title}{merged}{total}{data}.xlsx"
    wb = Workbook()
    ws = wb.active
    r = 1
    if title:
        ws.cell(1, 1, "Acme Corp - Roster")
        r = 3
    if merged:
        ws.cell(r, 1, "Employee")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        r += 1
    for j, h in enumerate(["Name", "SSN", "Zip"], start=1):
        ws.cell(r, j, h)
    rows = [
        ["Ann Lee", "123-45-6780", "10001"],
        ["Bo Diaz", "234-56-7890", "10002"],
        ["Cy Park", "345-67-8901", "10003"],
    ]
    for i, row in enumerate(rows if data else [], start=r + 1):
        for j, v in enumerate(row, start=1):
            ws.cell(i, j, v)
    if hidden and data:
        ws.row_dimensions[r + 2].hidden = True
    if total and data:
        ws.cell(r + 4, 1, "Total")
        ws.cell(r + 4, 2, 3)
    wb.save(p)
    return {
        c
        for s in read_file(p, "D1", p.name, "XLSX", S.sheets.reader)
        for c, _ in fix_layout(s, S.sheets.layout).codes
    }


def repair(fn: Callable[[], R.Fixed]) -> set[str]:
    return set(fn().codes)


RC = R.RepairContext(S.sheets.repair, S.zip_leading_zero_states, AS_OF)
T = S.thresholds
HDR = "Name,SSN,DOB,Zip,State,Card,Notes\n"


def P(code: str, fn: Callable[[], set[str]], note: str = "") -> Case:
    return Case(code, True, fn, note)


def N(code: str, fn: Callable[[], set[str]], note: str = "") -> Case:
    return Case(code, False, fn, note)


CASES: list[Case] = [
    # names
    P("NAME_MISSING", lambda: name("SMITH")),
    N("NAME_MISSING", lambda: name("JOHN SMITH")),
    P("NAME_MISSING", lambda: parts("", "", "SMITH"), "split columns, no first name"),
    P("NAME_TITLE_REMOVED", lambda: name("DR. KERRY M SMITH")),
    N("NAME_TITLE_REMOVED", lambda: name("KERRY M SMITH")),
    N("NAME_TITLE_REMOVED", lambda: name("DREW SMITH"), "DREW is a name, not DR"),
    P("NAME_SUFFIX_INVALID", lambda: name("Young JJ, Benjamin")),
    N("NAME_SUFFIX_INVALID", lambda: name("YOUNG JR, BENJAMIN")),
    N("NAME_SUFFIX_INVALID", lambda: name("BENJAMIN YOUNG Junior"), "variant normalized"),
    P("NAME_SUFFIX_IN_NAME", lambda: parts("JOHN", "", "SMITH JR")),
    N("NAME_SUFFIX_IN_NAME", lambda: parts("JOHN", "", "SMITH", "JR")),
    P("NAME_ORDER_AMBIGUOUS", lambda: name("Lee Allen", labelled=False)),
    N(
        "NAME_ORDER_AMBIGUOUS",
        lambda: name("Lee Allen", "FIRST_LAST", labelled=True),
        "field label fixes the order",
    ),
    N("NAME_ORDER_AMBIGUOUS", lambda: name("Allen, Lee", labelled=False), "comma rule"),
    P("NAME_PUNCT_FIXED", lambda: name("O' BRIEN, JOSHUA")),
    P("NAME_PUNCT_FIXED", lambda: name("KAREN GARCIA - LOPEZ")),
    N("NAME_PUNCT_FIXED", lambda: name("O'BRIEN, JOSHUA")),
    P("NAME_WHITESPACE", lambda: parts("  Susan ", "", "Williams  ")),
    P("NAME_WHITESPACE", lambda: name("JOHN  SMITH")),
    N("NAME_WHITESPACE", lambda: parts("Susan", "", "Williams")),
    P(
        "NAME_ORG_LIKE",
        lambda: excluded(
            classify_entity(
                "HARBORLINE LOGISTICS INC", load_entity_config(ROOT / "config" / "entities.yaml")
            )[0]
        ),
    ),
    N(
        "NAME_ORG_LIKE",
        lambda: (
            {"NAME_ORG_LIKE"}
            if classify_entity("MARY SMITH", load_entity_config(ROOT / "config" / "entities.yaml"))[
                0
            ]
            == ORG
            else set()
        ),
    ),
    P("NAME_VARIANT_IN_DOC", lambda: variants([("JOHN", "ANDERSON"), ("JOHN", "ANDERSEN")])),
    N("NAME_VARIANT_IN_DOC", lambda: variants([("JOHN", "ANDERSON"), ("JOHN", "ANDERSON")])),
    # addresses
    P("ADDR_MISSING", lambda: addr(city="TAMPA", state="FL", zip5="33602")),
    N(
        "ADDR_MISSING",
        lambda: addr(address_line1="315 E KENNEDY BLVD", city="TAMPA", state="FL", zip5="33602"),
    ),
    P("ADDR_LINE2_SPLIT", lambda: addr(address_line1="330 PARK BLVD UPPER")),
    P("ADDR_LINE2_SPLIT", lambda: addr(raw_address="12 MAIN ST APT 4\nTAMPA, FL 33602")),
    N(
        "ADDR_LINE2_SPLIT",
        lambda: addr(address_line1="6030 STEFANIE FLATS"),
        "FLATS is not FL + unit",
    ),
    P("ADDR_UNIT_FORMAT", lambda: addr(address_line1="501 N OREGON ST APT6F")),
    P("ADDR_UNIT_FORMAT", lambda: addr(address_line1="12 MAIN ST", address_line2="SUITE 200")),
    N("ADDR_UNIT_FORMAT", lambda: addr(address_line1="501 N OREGON ST APT 6F")),
    P("ADDR_STATE_INVALID", lambda: addr(address_line1="1 A ST", state="ZZ")),
    N("ADDR_STATE_INVALID", lambda: addr(address_line1="1 A ST", state="PR")),
    P(
        "ADDR_ZIP_FORMAT",
        lambda: addr(
            address_line1="1201 LEOPARD ST", city="CORPUS CHRISTI", state="TX", zip5="7840"
        ),
    ),
    P(
        "ADDR_ZIP_FORMAT",
        lambda: addr(address_line1="1 A ST", state="MA", zip5="2134"),
        "suggest zero-padding",
    ),
    N("ADDR_ZIP_FORMAT", lambda: addr(address_line1="1 A ST", state="MA", zip5="02134")),
    P(
        "ADDR_ZIP_STATE_MISMATCH",
        lambda: addr(address_line1="1 A ST", city="SACRAMENTO", state="NV", zip5="95814"),
    ),
    N(
        "ADDR_ZIP_STATE_MISMATCH",
        lambda: addr(address_line1="1 A ST", city="SACRAMENTO", state="CA", zip5="95814"),
    ),
    P(
        "ADDR_ZIP_CITY_MISMATCH",
        lambda: addr(address_line1="828 I ST", city="SACRAMNETO", state="CA", zip5="95814"),
    ),
    N(
        "ADDR_ZIP_CITY_MISMATCH",
        lambda: addr(address_line1="828 I ST", city="SACRAMENTO", state="CA", zip5="95814"),
    ),
    P("ADDR_PO_BOX", lambda: addr(address_line1="PO BOX 123")),
    P("ADDR_PO_BOX", lambda: addr(address_line1="P.O. BOX 77")),
    N("ADDR_PO_BOX", lambda: addr(address_line1="123 BOXWOOD LN")),
    # dates of birth
    P("DOB_UNPARSEABLE", lambda: dob(None, "31/31/1990")),
    N("DOB_UNPARSEABLE", lambda: dob("1990-01-31", "01/31/1990")),
    N("DOB_UNPARSEABLE", lambda: dob(None, None), "no DOB at all"),
    P("DOB_TWO_DIGIT_YEAR", lambda: dob("1982-06-23", "06/23/82")),
    N("DOB_TWO_DIGIT_YEAR", lambda: dob("1982-06-23", "06/23/1982")),
    P("DOB_FUTURE", lambda: dob("2027-01-01")),
    N("DOB_FUTURE", lambda: dob("2026-10-01"), "as_of_date itself"),
    P("DOB_IMPLAUSIBLE", lambda: dob("1900-01-01")),
    N("DOB_IMPLAUSIBLE", lambda: dob("1907-01-01"), "119 years"),
    P("DOB_MINOR", lambda: dob("2010-01-01")),
    N("DOB_MINOR", lambda: dob("2008-10-01"), "18 on as_of_date"),
    # identifiers
    P("SSN_FORMAT", lambda: ssn("12345678")),
    N("SSN_FORMAT", lambda: ssn("123-45-6780")),
    P("SSN_AREA_INVALID", lambda: ssn("666-12-3456")),
    P("SSN_AREA_INVALID", lambda: ssn("000-12-3456")),
    P("SSN_AREA_INVALID", lambda: ssn("950-12-3456"), "9xx outside the ITIN ranges"),
    N("SSN_AREA_INVALID", lambda: ssn("665-12-3456")),
    N("SSN_AREA_INVALID", lambda: ssn("900-70-1234"), "ITIN wins (ADR-035)"),
    P("SSN_GROUP_INVALID", lambda: ssn("123-00-4567")),
    N("SSN_GROUP_INVALID", lambda: ssn("123-01-4567")),
    P("SSN_SERIAL_INVALID", lambda: ssn("123-45-0000")),
    N("SSN_SERIAL_INVALID", lambda: ssn("123-45-0001")),
    P("SSN_KNOWN_INVALID", lambda: ssn("078-05-1120")),
    N("SSN_KNOWN_INVALID", lambda: ssn("078-05-1121")),
    P("ITIN_DETECTED", lambda: ssn("900-70-1234")),
    N("ITIN_DETECTED", lambda: ssn("900-12-1234"), "group 12 not ITIN"),
    P("EIN_EXCLUDED", lambda: excluded("EIN")),
    P("EIN_EXCLUDED", lambda: excluded("PROVIDER_TIN")),
    N("EIN_EXCLUDED", lambda: excluded("NPI")),
    P("NPI_EXCLUDED", lambda: excluded("NPI")),
    N("NPI_EXCLUDED", lambda: excluded("EIN")),
    P("DL_FORMAT_STATE", lambda: set(dl_codes("12345", "CA", V.dl_regex))),
    N("DL_FORMAT_STATE", lambda: set(dl_codes("A1234567", "CA", V.dl_regex))),
    N(
        "DL_FORMAT_STATE",
        lambda: set(dl_codes("12345", "WY", V.dl_regex)),
        "state without configured formats",
    ),
    P("DL_STATE_MISSING", lambda: set(dl_codes("A1234567", None, V.dl_regex))),
    N("DL_STATE_MISSING", lambda: set(dl_codes(None, None, V.dl_regex))),
    P("PASSPORT_FORMAT", lambda: set(passport_codes("AB1234567"))),
    N("PASSPORT_FORMAT", lambda: set(passport_codes("A12345678"))),
    N("PASSPORT_FORMAT", lambda: set(passport_codes("123456789"))),
    P("CARD_LUHN_FAIL", lambda: set(card_codes("4111111111111112", "TEXT"))),
    N("CARD_LUHN_FAIL", lambda: set(card_codes("4111111111111111", "TEXT"))),
    N(
        "CARD_LUHN_FAIL",
        lambda: set(card_codes("************1112", "PREMASKED")),
        "pre-masked is not checked",
    ),
    P("ROUTING_CHECKSUM_FAIL", lambda: set(routing_codes("021000022"))),
    N("ROUTING_CHECKSUM_FAIL", lambda: set(routing_codes("021000021"))),
    # spreadsheet producers (M3)
    P("SHEET_HEADER_OFFSET", lambda: xlsx_layout(title=True)),
    N("SHEET_HEADER_OFFSET", lambda: xlsx_layout()),
    P("SHEET_MERGED_HEADER", lambda: xlsx_layout(merged=True)),
    N("SHEET_MERGED_HEADER", lambda: xlsx_layout()),
    P("SHEET_HIDDEN_ROWS", lambda: xlsx_layout(hidden=True)),
    N("SHEET_HIDDEN_ROWS", lambda: xlsx_layout()),
    P("SHEET_TOTAL_ROW_DROPPED", lambda: xlsx_layout(total=True)),
    N("SHEET_TOTAL_ROW_DROPPED", lambda: xlsx_layout()),
    P("SHEET_SKIPPED", lambda: xlsx_layout(data=False)),
    N("SHEET_SKIPPED", lambda: xlsx_layout()),
    P(
        "CSV_ENCODING_DETECTED",
        lambda: sheet("Name,Zip\nJosé Muñoz,10001\nAna Peña,10002\n", "cp1252"),
    ),
    N("CSV_ENCODING_DETECTED", lambda: sheet("Name,Zip\nJosé Muñoz,10001\nAna Peña,10002\n")),
    P("COL_MAPPING_LOW_CONF", lambda: set(decide(None, "DOB", 0.65, T).codes)),
    N("COL_MAPPING_LOW_CONF", lambda: set(decide(None, "DOB", 0.75, T).codes)),
    P("COL_MAPPING_AGREE_LOW_PROBA", lambda: set(decide("SSN", "SSN", 0.4, T).codes)),
    N("COL_MAPPING_AGREE_LOW_PROBA", lambda: set(decide("SSN", "SSN", 0.9, T).codes)),
    P("COL_MAPPING_ML_OVERRIDE", lambda: set(decide("PHONE", "SSN", 0.85, T).codes)),
    N("COL_MAPPING_ML_OVERRIDE", lambda: set(decide("PHONE", "SSN", 0.75, T).codes)),
    P(
        "COL_UNMAPPED_SENSITIVE",
        lambda: (
            {"COL_UNMAPPED_SENSITIVE"} if unmapped_sensitive("OTHER", {"ssn_valid": 0.9}) else set()
        ),
    ),
    N(
        "COL_UNMAPPED_SENSITIVE",
        lambda: (
            {"COL_UNMAPPED_SENSITIVE"} if unmapped_sensitive("SSN", {"ssn_valid": 0.9}) else set()
        ),
    ),
    P(
        "XL_ZIP_LEADING_ZERO_RESTORED",
        lambda: sheet(HDR + "Ann Lee,123-45-6780,,2134,MA,,\nBo Diaz,234-56-7890,,2135,MA,,\n"),
    ),
    N(
        "XL_ZIP_LEADING_ZERO_RESTORED",
        lambda: sheet(HDR + "Ann Lee,123-45-6780,,02134,MA,,\nBo Diaz,234-56-7890,,02135,MA,,\n"),
    ),
    P("XL_SSN_ZERO_PADDED", lambda: repair(lambda: R.repair_ssn(12345678, RC))),
    N("XL_SSN_ZERO_PADDED", lambda: repair(lambda: R.repair_ssn("012-34-5678", RC))),
    P(
        "XL_DATE_SERIAL_CONVERTED",
        lambda: sheet(
            HDR
            + "Ann Lee,123-45-6780,32874,10001,NY,,\nBo Diaz,234-56-7890,1990-01-02,10002,NY,,\n"
        ),
    ),
    N(
        "XL_DATE_SERIAL_CONVERTED",
        lambda: sheet(
            HDR + "Ann Lee,123-45-6780,1990-01-01,10001,NY,,\n"
            "Bo Diaz,234-56-7890,1990-01-02,10002,NY,,\n"
        ),
    ),
    P("DOB_AMBIGUOUS_FORMAT", lambda: repair(lambda: R.repair_date("03/04/1985", RC))),
    N(
        "DOB_AMBIGUOUS_FORMAT",
        lambda: repair(
            lambda: R.repair_date(
                "03/04/1985", R.RepairContext(RC.cfg, RC.zip_states, AS_OF, False)
            )
        ),
        "column evidence decides",
    ),
    P(
        "XL_NUMBER_PRECISION_LOST",
        lambda: sheet(
            HDR + "Ann Lee,123-45-6780,,10001,NY,4.11111E+15,\n"
            "Bo Diaz,234-56-7890,,10002,NY,4111111111111111,\n"
        ),
    ),
    N(
        "XL_NUMBER_PRECISION_LOST",
        lambda: sheet(
            HDR
            + "Ann Lee,123-45-6780,,10001,NY,4111111111111111,\nBo Diaz,234-56-7890,,10002,NY,,\n"
        ),
    ),
    P(
        "XL_ERROR_VALUE",
        lambda: sheet(
            HDR + "Ann Lee,123-45-6780,#N/A,10001,NY,,\nBo Diaz,234-56-7890,,10002,NY,,\n"
        ),
    ),
    N(
        "XL_ERROR_VALUE",
        lambda: sheet(
            HDR + "Ann Lee,123-45-6780,N/A text,10001,NY,,\nBo Diaz,234-56-7890,,10002,NY,,\n"
        ),
    ),
    P(
        "PREMASKED_VALUE",
        lambda: sheet(HDR + "Ann Lee,XXX-XX-6780,,10001,NY,,\nBo Diaz,234-56-7890,,10002,NY,,\n"),
    ),
    N(
        "PREMASKED_VALUE",
        lambda: sheet(HDR + "Ann Lee,123-45-6780,,10001,NY,,\nBo Diaz,234-56-7890,,10002,NY,,\n"),
    ),
    P(
        "FREE_TEXT_PII_FOUND",
        lambda: sheet(
            HDR
            + "Ann Lee,,,10001,NY,,Verified SSN 123-45-6780\nBo Diaz,234-56-7890,,10002,NY,,ok\n"
        ),
    ),
    N(
        "FREE_TEXT_PII_FOUND",
        lambda: sheet(
            HDR + "Ann Lee,,,10001,NY,,Seen 3 times\nBo Diaz,234-56-7890,,10002,NY,,ok\n"
        ),
    ),
    P(
        "ROW_DUPLICATE_IN_FILE",
        lambda: sheet(
            HDR + "Ann Lee,123-45-6780,,10001,NY,,\nAnn Lee,123-45-6780,,10001,NY,,\n"
            "Bo Diaz,234-56-7890,,10002,NY,,\n"
        ),
    ),
    N(
        "ROW_DUPLICATE_IN_FILE",
        lambda: sheet(HDR + "Ann Lee,123-45-6780,,10001,NY,,\nBo Diaz,234-56-7890,,10002,NY,,\n"),
    ),
    # completeness and person level
    P(
        "PERSON_NO_QUALIFYING_ELEMENT",
        lambda: set(
            C.qualifying_codes(
                {"raw_name": "MELISSA ESPINOZA", "phone": "+12125550101"},
                V.rules.qualifying_elements,
            )
        ),
    ),
    N(
        "PERSON_NO_QUALIFYING_ELEMENT",
        lambda: set(C.qualifying_codes({"dob": "1990-01-01"}, V.rules.qualifying_elements)),
    ),
    N(
        "PERSON_NO_QUALIFYING_ELEMENT",
        lambda: set(
            C.qualifying_codes({"_raw_card": "************1111"}, V.rules.qualifying_elements)
        ),
        "pre-masked card is present",
    ),
    P(
        "LOW_OCR_CONFIDENCE",
        lambda: set(
            C.low_ocr_codes(
                {"extraction_method": "OCR", "field_confidence": json.dumps({"ssn": 0.62})}, 0.8
            )
        ),
    ),
    N(
        "LOW_OCR_CONFIDENCE",
        lambda: set(
            C.low_ocr_codes(
                {"extraction_method": "OCR", "field_confidence": json.dumps({"ssn": 0.93})}, 0.8
            )
        ),
    ),
    N(
        "LOW_OCR_CONFIDENCE",
        lambda: set(
            C.low_ocr_codes(
                {"extraction_method": "SHEET", "field_confidence": json.dumps({"SSN": 0.5})}, 0.8
            )
        ),
        "mapping proba",
    ),
    P("DOC_TYPE_UNKNOWN", lambda: doc_type("UNKNOWN")),
    N("DOC_TYPE_UNKNOWN", lambda: doc_type("HEALTH_CLAIM")),
    P("FERPA_RECORD", lambda: set(C.ferpa_codes({"_raw_student_id": "S1234567"}))),
    N("FERPA_RECORD", lambda: set(C.ferpa_codes({"employer_name": "ACME"}))),
]


def counts() -> dict[str, Counter]:
    out: dict[str, Counter] = {}
    for c in CASES:
        out.setdefault(c.code, Counter())["pos" if c.positive else "neg"] += 1
    return out
