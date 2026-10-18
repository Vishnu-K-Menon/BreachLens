from pathlib import Path

from openpyxl import Workbook

from breachlens.sheets.layout import fix_layout
from breachlens.sheets.profile import column_profiles, file_profile, shape
from breachlens.sheets.reader import as_text, detect_encoding, read_file
from breachlens.sheets.settings import load_sheet_settings

ROOT = Path(__file__).resolve().parents[2]
S = load_sheet_settings(ROOT / "config")


def _roster(path: Path) -> None:
    """Title rows, a merged 2-row header, a hidden row and a total row (P27-P30)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Roster"
    ws["A1"] = "Acme Corp - Employee Roster"
    ws["A2"] = "Exported 2026-09-30"
    ws["A4"], ws["C4"] = "Employee", "Address"
    ws.merge_cells("A4:B4")
    ws.merge_cells("C4:D4")
    for col, h in zip("ABCD", ["Name", "SSN", "Street", "Zip"], strict=True):
        ws[f"{col}5"] = h
    data = [
        ["Ann Lee", "123-45-6789", "1 Main St", "02134"],
        ["Bo Diaz", "234-56-7890", "2 Oak Ave", "10001"],
        ["Cy Park", "345-67-8901", "3 Elm Rd", "73301"],
    ]
    for i, row in enumerate(data, start=6):
        for col, v in zip("ABCD", row, strict=True):
            ws[f"{col}{i}"] = v
    ws.row_dimensions[7].hidden = True
    ws["A9"], ws["B9"] = "Total", 3
    wb.save(path)


def test_xlsx_layout_repairs(tmp_path):
    p = tmp_path / "roster.xlsx"
    _roster(p)
    [sheet] = read_file(p, "D0001", "roster.xlsx", "XLSX", S.sheets.reader)
    t = fix_layout(sheet, S.sheets.layout)
    assert t.header_row == 5
    assert t.headers == ["Employee Name", "Employee SSN", "Address Street", "Address Zip"]
    assert [r for r, _ in t.rows] == [6, 7, 8]
    assert t.total_rows_dropped == 1 and t.hidden_data_rows == [7]
    codes = {c for c, _ in t.codes}
    assert {
        "SHEET_MERGED_HEADER",
        "SHEET_HEADER_OFFSET",
        "SHEET_TOTAL_ROW_DROPPED",
        "SHEET_HIDDEN_ROWS",
    } <= codes
    prof = file_profile(t)
    assert (prof["n_rows"], prof["hidden_rows"], prof["skipped"]) == (3, 1, False)


def test_column_profile_keeps_shapes_not_values(tmp_path):
    p = tmp_path / "roster.xlsx"
    _roster(p)
    [sheet] = read_file(p, "D0001", "roster.xlsx", "XLSX", S.sheets.reader)
    cols = column_profiles(fix_layout(sheet, S.sheets.layout))
    ssn = next(c for c in cols if c["header_raw"] == "Employee SSN")
    assert ssn["top_shapes"] == "ddd-dd-dddd (3)" and ssn["null_pct"] == 0.0
    assert "123-45-6789" not in str(cols)


def test_shape_masks_characters():
    assert shape("Ab-12 x") == "Aa-dd a"
    assert shape("1234567890123456") == "dddddddddddddd…"


def test_csv_encoding_delimiter_and_comment(tmp_path):
    p = tmp_path / "pay.csv"
    text = (
        "# SYNTHETIC - NOT REAL\nName;SS#;Zip\n"
        "José Muñoz;123-45-6789;02134\nAna Peña;234-56-7890;10001\n"
    )
    p.write_bytes(text.encode("cp1252"))
    [sheet] = read_file(p, "D0002", "pay.csv", "CSV", S.sheets.reader)
    assert sheet.encoding == "cp1252" and sheet.delimiter == ";"
    t = fix_layout(sheet, S.sheets.layout)
    assert t.headers == ["Name", "SS#", "Zip"]
    assert as_text(t.rows[0][1][0]) == "José Muñoz"
    assert as_text(t.rows[0][1][2]) == "02134"  # CSV text keeps the leading zero


def test_detect_encoding_prefers_utf8_then_cp1252():
    assert detect_encoding("Muñoz".encode(), ["utf-8", "cp1252", "latin-1"]) == "utf-8"
    assert detect_encoding("Muñoz".encode("cp1252"), ["utf-8", "cp1252", "latin-1"]) == "cp1252"


def test_as_text():
    assert as_text(None) == "" and as_text(12.0) == "12" and as_text(1.5) == "1.5"
