"""LibreOffice recalculation of the review workbook (M7, ADR-041/068)."""

from datetime import date
from pathlib import Path

import pytest
from run_fixture import write_run

from breachlens.config import LeakGuardConfig
from breachlens.db.load import build_database
from breachlens.privacy.leak_guard import LeakGuard
from breachlens.report.excel_writer import load_report_data, write_workbook
from breachlens.report.recalc import check, recalculate
from breachlens.validate.settings import load_validation_settings

ROOT = Path(__file__).resolve().parents[2]
GUARD = LeakGuard(LeakGuardConfig())


@pytest.mark.libreoffice
def test_recalculated_workbook_has_no_errors_and_matches_python(tmp_path):
    run_dir = write_run(tmp_path)
    db = tmp_path / "output" / "breachlens.duckdb"
    build_database(run_dir, db, ROOT, GUARD)
    d = load_report_data(db, ROOT, date(2026, 10, 1))
    path = write_workbook(d, run_dir, GUARD, load_validation_settings(ROOT / "config").severity)
    res = check(recalculate(path, tmp_path / "recalc"), d)
    assert res.errors == []
    assert res.mismatches == []
    assert len(res.summary) >= 20
    assert (res.status_rows, res.status_disagree) == (2, 0)
    assert res.ok


def test_check_reports_a_wrong_summary_value(tmp_path):
    """Without LibreOffice: a hand-made 'recalculated' file with one wrong count."""
    from openpyxl import Workbook

    from breachlens.report.excel_writer import SUMMARY_HEADER_ROW, summary_rows

    run_dir = write_run(tmp_path)
    db = tmp_path / "output" / "breachlens.duckdb"
    build_database(run_dir, db, ROOT, GUARD)
    d = load_report_data(db, ROOT, date(2026, 10, 1))
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    for k, r in enumerate(summary_rows(d)):
        value = r.python + 1 if r.item == "People" else r.python
        ws.cell(SUMMARY_HEADER_ROW + 2 + k, 1, r.section)
        ws.cell(SUMMARY_HEADER_ROW + 2 + k, 2, r.item)
        ws.cell(SUMMARY_HEADER_ROW + 2 + k, 3, value)
    nl = wb.create_sheet("Notification_List")
    nl.append(["person_id", "status_agrees"])
    nl.append(["P0001", True])
    nl.append(["P0002", "#N/A"])
    path = tmp_path / "fake.xlsx"
    wb.save(path)
    res = check(path, d)
    assert [m[1] for m in res.mismatches] == ["People"]
    assert res.errors == [("Notification_List", "B3", "#N/A")]
    assert res.status_disagree == 1 and not res.ok
