"""Excel review workbook (M7, ADR-068) on the tiny hand-built run of the M6 tests."""

from datetime import date
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook
from run_fixture import RUN, write_run

from breachlens.config import LeakGuardConfig
from breachlens.db.load import build_database
from breachlens.privacy.leak_guard import LeakError, LeakGuard, scan_dir
from breachlens.report.excel_writer import (
    SHEETS,
    TABLES,
    load_report_data,
    summary_rows,
    write_workbook,
)
from breachlens.validate.settings import load_validation_settings

ROOT = Path(__file__).resolve().parents[2]
GUARD = LeakGuard(LeakGuardConfig())
SEVERITY = load_validation_settings(ROOT / "config").severity


@pytest.fixture
def data(tmp_path):
    run_dir = write_run(tmp_path)
    pd.DataFrame(
        [
            {
                "doc_id": "D0001",
                "page": "",
                "kind": "ORG",
                "value_masked": "ACME SUPPLY LLC",
                "reason": "NAME_ORG_LIKE",
                "sheet": "S",
                "row": "7.0",
            }
        ]
    ).to_csv(run_dir / "excluded_entities.csv", index=False)
    db = tmp_path / "output" / "breachlens.duckdb"
    build_database(run_dir, db, ROOT, GUARD)
    return load_report_data(db, ROOT, date(2026, 10, 1)), run_dir


def test_workbook_sheets_tables_names_and_formulas(data, tmp_path):
    d, run_dir = data
    path = write_workbook(d, run_dir, GUARD, SEVERITY, vba=tmp_path / "missing.bin")
    assert path.name == f"breachlens_{RUN}.xlsx"  # no vbaProject.bin -> no macros
    wb = load_workbook(path)
    assert tuple(wb.sheetnames) == SHEETS
    tables = {name: ws.tables for name, ws in ((n, wb[n]) for n in TABLES)}
    assert {sheet: list(t) for sheet, t in tables.items()} == {s: [t] for s, t in TABLES.items()}
    assert set(wb.defined_names) >= {"RunDate", "RunId"}
    assert wb.defined_names["RunDate"].attr_text == "Summary!$B$3"
    assert wb["Summary"]["B3"].value.date() == date(2026, 10, 1)
    formulas = [c.value for row in wb["Summary"].iter_rows(min_col=3, max_col=3) for c in row]
    assert "=COUNTIFS(tblPeople[notify_required],TRUE)" in formulas
    assert any(str(f).startswith("=SUMPRODUCT((tblPeople[ssn_exposed]=TRUE)") for f in formulas)
    notify = wb["Notification_List"]
    header = [c.value for c in notify[1]]
    status = notify.cell(2, header.index("status_xl") + 1).value
    status = getattr(status, "text", status)  # XLOOKUP may come back as an array formula
    assert "XLOOKUP" in status and "IF(AND(" in status
    assert "tblNotify[[#This Row],[notify_required]]" in status  # qualified for LibreOffice
    people = wb["People"]
    ph = [c.value for c in people[1]]
    assert "DATEDIF" in people.cell(2, ph.index("age_at_rundate") + 1).value
    assert people.cell(2, ph.index("zip5") + 1).value == "02134"  # leading zero kept
    excluded = wb["Excluded"]
    assert [c.value for c in excluded[2]][:4] == [
        "entities",
        "ORG",
        "ACME SUPPLY LLC",
        "NAME_ORG_LIKE",
    ]


def test_review_queue_one_row_per_column_and_resolved(data):
    d, run_dir = data
    q = d.review
    cols = q[q["item_type"] == "COLUMN"]
    assert len(cols) == 2  # ID (open) + Insurer (reviewed), once each
    assert cols.set_index("header")["review_status"].to_dict() == {
        "ID": "PENDING",
        "Insurer": "RESOLVED",
    }
    assert cols.set_index("header")["affected_rows"].to_dict() == {"ID": 1, "Insurer": 1}
    assert list(q["status"]).index("RESOLVED") > list(q["status"]).index("OPEN")  # open first
    path = write_workbook(d, run_dir, GUARD, SEVERITY)
    ws = load_workbook(path)["Review_Queue"]
    header = [c.value for c in ws[4]]
    assert {"review_status", "reviewer_note", "reviewed_by", "reviewed_at", "go"} <= set(header)
    dv = [v for v in ws.data_validations.dataValidation if v.type == "list"]
    assert dv and "PENDING" in dv[0].formula1


def test_no_key_or_raw_columns_and_guard_scans_the_workbook(data):
    d, run_dir = data
    path = write_workbook(d, run_dir, GUARD, SEVERITY)
    wb = load_workbook(path)
    for name in TABLES:
        header = [c.value for c in wb[name][4 if name == "Review_Queue" else 1]]
        assert not [h for h in header if h and (h.endswith("_key") or h.startswith("_"))]
        assert "raw_name" not in header and "run_id" not in header
    files, findings = scan_dir(run_dir, GUARD)
    assert path in files and findings == []


def test_leak_guard_blocks_the_workbook_and_saves_nothing(data):
    d, run_dir = data
    d.notify.loc[0, "notify_note"] = "call 123-45-6789"
    with pytest.raises(LeakError):
        write_workbook(d, run_dir, GUARD, SEVERITY)
    assert not list(run_dir.glob("*.xls*")) and not list(run_dir.glob(".tmp_*"))


def test_summary_rows_python_side(data):
    d, _ = data
    rows = {(r.section, r.item): r.python for r in summary_rows(d)}
    assert rows[("KPI", "People")] == 2
    assert rows[("KPI", "People to notify")] == 1
    assert rows[("KPI", "Minors (flag)")] == 1
    assert rows[("KPI", "Minors (DATEDIF from RunDate)")] == 1
    assert rows[("KPI", "Open COLUMN items")] == 1
    assert rows[("KPI", "Resolved COLUMN items (review_overrides.yaml)")] == 1
    assert rows[("To notify by state", "MA")] == 1
    assert rows[("Top reason codes (mentions)", "SSN_AREA_INVALID")] == 1


def test_macro_module_is_present_and_self_contained():
    bas = (ROOT / "excel" / "ReviewHelpers.bas").read_bytes()
    text = bas.decode("utf-8")
    assert b"\r\n" in bas and "Option Explicit" in text  # VBA import expects CRLF
    for macro in ("Public Sub ShowPendingOnly()", "Public Sub ExportDecisions()"):
        assert macro in text
    assert "On Error GoTo" in text and "tblReview" in text
    for forbidden in ("CreateObject", "XMLHTTP", "WinHttp", "Shell", "Declare "):
        assert forbidden not in text
